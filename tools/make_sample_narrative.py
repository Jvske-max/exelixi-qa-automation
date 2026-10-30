"""Generate synthetic narrative test case document for test-normalizer evaluations."""

from __future__ import annotations

import argparse
import datetime
import sys
from pathlib import Path

import docx

DEFAULT_OUTPUT_PATH = Path("build/sample/muestra.docx")


def create_sample_narrative(out_path: Path | str = DEFAULT_OUTPUT_PATH) -> Path:
    """Generate synthetic narrative document with two test cases (one complete, one with gaps)."""
    target = Path(out_path)
    target.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()

    # a. Tabla "Identificación"
    headers = [
        "Fecha",
        "Módulo",
        "Identificador del caso",
        "Descripción corta",
        "Descripción completa",
        "Ambiente",
        "Analista Programador",
        "Analista de prueba",
    ]
    values = [
        datetime.date.today().isoformat(),
        "Usuarios",
        "R-SAMPLE-01",
        "Prueba sintética de usuarios",
        "Prueba sintética narrativa para normalizador",
        "QA",
        "Dev Test",
        "QA Test",
    ]

    table = doc.add_table(rows=2, cols=len(headers))
    for col_idx, (h, v) in enumerate(zip(headers, values)):
        table.rows[0].cells[col_idx].text = h
        table.rows[1].cells[col_idx].text = v

    # b. Caso 1: Creación de usuario regular (completo y válido)
    doc.add_paragraph("Link: https://example.com/portal")
    doc.add_paragraph("Caso 1: Creación de usuario regular")
    doc.add_paragraph("Objetivo de la Prueba: Verificar la creación de usuario en el portal")
    doc.add_paragraph("Iniciamos sesión en el sistema con el usuario ${QA_USER_ADMIN} y clave ${QA_PASS_ADMIN}.")
    doc.add_paragraph("Ingresamos al link https://example.com/portal")
    doc.add_paragraph("Escribimos ${QA_USER_ADMIN} en el campo Usuario")
    doc.add_paragraph("Pulsamos el botón Ingresar")
    doc.add_paragraph("Verificamos Bienvenido al Sistema")

    # c. Caso 2: Aprobación especial con segundo factor (sin link, verificación vaga, flujo OTP)
    doc.add_paragraph("Caso 2: Aprobación especial con segundo factor")
    doc.add_paragraph("Objetivo de la Prueba: Aprobación de solicitud")
    doc.add_paragraph("El sistema funciona correctamente.")
    doc.add_paragraph("Se solicita el código OTP enviado al teléfono móvil.")

    doc.save(str(target))
    return target


def main(argv: list[str] | None = None) -> int:
    """CLI entrypoint for generating synthetic sample narrative document."""
    parser = argparse.ArgumentParser(description="Generate synthetic sample narrative .docx file")
    parser.add_argument(
        "--out",
        type=Path,
        default=DEFAULT_OUTPUT_PATH,
        help=f"Output file path (default: {DEFAULT_OUTPUT_PATH})",
    )
    args = parser.parse_args(argv)

    created_path = create_sample_narrative(args.out)
    print(f"Sample narrative created successfully: {created_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
