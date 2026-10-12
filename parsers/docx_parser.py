"""Word (.docx) specification parser extracting metadata and steps."""

from __future__ import annotations

import re
from pathlib import Path

import docx

from parsers.models import StepDefinition, TestCase

PSR_STEP_PATTERN = re.compile(
    r"^Step\s+(\d+):\s*(?:\(([^)]+)\)\s*)?User\s+([a-zA-Z\s]+?)\s+(?:on|in)\s+\"([^\"]+)\"(?:\s+in\s+\"([^\"]*)\")?(?:\s*\[(.*)\])?$",
    re.IGNORECASE,
)

STD_STEP_PATTERN = re.compile(
    r"^(?:Step|Paso)\s+(\d+)[:.]\s*(.*)$",
    re.IGNORECASE,
)

NOISE_STEP_PATTERN = re.compile(
    r"^Step\s+\d+\s+screenshot\.?$",
    re.IGNORECASE,
)


def normalize_psr_action(raw_action: str) -> str:
    """Normalize PSR action string to standard action verb."""
    act = raw_action.lower().strip()
    if "click" in act:
        return "double_click" if "double" in act else "click"
    if "keyboard" in act or "input" in act:
        return "fill"
    if "wheel" in act or "scroll" in act:
        return "scroll"
    if "drag" in act:
        return "drag"
    return act


def parse_standard_description(desc: str) -> tuple[str, str, str | None]:
    """Parse a standard step description into action, target, and optional value."""
    text = desc.strip()
    # Check navigation
    m_nav = re.search(r"(?:navigate\s+to|navegar\s+a|ir\s+a|abrir)\s+[\"']?(\S+?)[\"']?$", text, re.IGNORECASE)
    if m_nav:
        return "navigate", m_nav.group(1), None

    # Check fill: Escribir <val> en <target> or Fill <val> in <target>
    m_fill = re.search(
        r"(?:fill|escribir|type|ingresar)\s+[\"']?([^\"']+)[\"']?\s+(?:in|into|en)\s+[\"']?([^\"']+)[\"']?$",
        text,
        re.IGNORECASE,
    )
    if m_fill:
        return "fill", m_fill.group(2).strip(), m_fill.group(1).strip()

    # Check click
    m_click = re.search(r"(?:click|clic|presionar)\s+(?:on|en)?\s*[\"']?([^\"']+)[\"']?$", text, re.IGNORECASE)
    if m_click:
        return "click", m_click.group(1).strip(), None

    # Check verify
    m_ver = re.search(r"(?:verify|verificar|validar|assert)\s+[\"']?([^\"']+)[\"']?$", text, re.IGNORECASE)
    if m_ver:
        return "verify", m_ver.group(1).strip(), None

    return "click", text, None


class DocxParser:
    """Parser for Microsoft Word (.docx) test specifications."""

    @classmethod
    def parse(cls, file_path: str | Path) -> TestCase:
        """Parse a .docx document into a TestCase model."""
        path = Path(file_path)
        doc = docx.Document(str(path))

        table_data = cls._extract_table_data(doc)
        case_id = table_data.get("identificador del caso") or table_data.get("caso") or ""
        module = table_data.get("módulo") or table_data.get("modulo") or ""
        environment = table_data.get("ambiente de prueba") or table_data.get("ambiente") or "QA"
        developer = (
            table_data.get("analista programador")
            or table_data.get("analista de prueba")
            or table_data.get("analista")
            or "Exelixi QA"
        )
        title_from_table = table_data.get("descripción completa") or table_data.get("descripción corta") or ""

        # Extract metadata from paragraphs
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        url = None
        title = ""
        objective = ""
        credentials: dict[str, str] = {}

        for i, text in enumerate(paragraphs):
            # Check URL
            if not url:
                m_url = re.search(r"(?:Link|URL):\s*(https?://\S+)", text, re.IGNORECASE)
                if m_url:
                    url = m_url.group(1).strip()

            # Check title / Caso
            if not title:
                m_title = re.search(r"(?:Caso\s*#?\s*\d*:?\s*)(.+)", text, re.IGNORECASE)
                if m_title:
                    title = text.strip()
                else:
                    m_title_alt = re.search(r"(?:Título|Titulo|Title):\s*(.+)", text, re.IGNORECASE)
                    if m_title_alt:
                        title = m_title_alt.group(1).strip()

            # Check objective
            if not objective:
                m_obj = re.search(r"^Objetivo(?:\s+de\s+la\s+Prueba)?\s*:\s*(.+)", text, re.IGNORECASE)
                if m_obj:
                    objective = m_obj.group(1).strip()
                elif re.match(r"^Objetivo(?:\s+de\s+la\s+Prueba)?\s*$", text, re.IGNORECASE) and i + 1 < len(paragraphs):
                    objective = paragraphs[i + 1].strip()

            # Check credentials
            m_user = re.search(r"(?:Usuario|User|Username):\s*(\S+)", text, re.IGNORECASE)
            if m_user:
                credentials["username"] = m_user.group(1).strip()
                credentials["usuario"] = m_user.group(1).strip()

            m_pass = re.search(r"(?:Clave|Password|Contraseña|Contrasena):\s*(\S+)", text, re.IGNORECASE)
            if m_pass:
                credentials["password"] = m_pass.group(1).strip()
                credentials["clave"] = m_pass.group(1).strip()

        # Fallbacks for metadata
        if not title:
            title = title_from_table or path.stem
        if not case_id:
            m_id = re.search(r"(R-\d+)", path.name)
            case_id = m_id.group(1) if m_id else "UNKNOWN"
        if not module:
            module = "General"

        # Extract steps
        steps = cls._extract_steps(paragraphs)

        return TestCase(
            case_id=case_id,
            module=module,
            title=title,
            objective=objective,
            credentials=credentials,
            steps=steps,
            source_file=str(path),
            environment=environment,
            developer=developer,
            url=url,
        )

    @classmethod
    def _extract_table_data(cls, doc: docx.Document) -> dict[str, str]:
        """Extract key-value metadata from Word document tables."""
        data: dict[str, str] = {}
        for table in doc.tables:
            rows = table.rows
            if len(rows) >= 2:
                headers = [c.text.strip().lower() for c in rows[0].cells]
                values = [c.text.strip() for c in rows[1].cells]
                for h, v in zip(headers, values):
                    if h and v and h != v.lower():
                        data[h] = v

            for row in rows:
                cells = [c.text.strip() for c in row.cells]
                if len(cells) >= 2 and cells[0]:
                    k = cells[0].lower()
                    v = cells[1].strip()
                    if v and k != v.lower():
                        data[k] = v
        return data

    @classmethod
    def _extract_steps(cls, paragraphs: list[str]) -> list[StepDefinition]:
        """Parse step definitions from document paragraphs."""
        steps: list[StepDefinition] = []

        for raw_text in paragraphs:
            # Clean invisible unicode marks common in Windows PSR
            clean = raw_text.replace("\u200e", "").replace("\u200f", "").strip()

            # Ignore noise lines
            if not clean or NOISE_STEP_PATTERN.match(clean) or clean.lower() == "previous next":
                continue

            # Try PSR pattern
            m_psr = PSR_STEP_PATTERN.match(clean)
            if m_psr:
                step_num = int(m_psr.group(1))
                timestamp = m_psr.group(2).strip() if m_psr.group(2) else None
                action = normalize_psr_action(m_psr.group(3))
                target = m_psr.group(4).strip()
                value = m_psr.group(6).strip() if m_psr.group(6) else None
                steps.append(
                    StepDefinition(
                        step_num=step_num,
                        raw_text=clean,
                        action=action,
                        target=target,
                        value=value,
                        timestamp=timestamp,
                    )
                )
                continue

            # Try standard numbered step pattern
            m_std = STD_STEP_PATTERN.match(clean)
            if m_std:
                step_num = int(m_std.group(1))
                desc = m_std.group(2).strip()
                action, target, value = parse_standard_description(desc)
                steps.append(
                    StepDefinition(
                        step_num=step_num,
                        raw_text=clean,
                        action=action,
                        target=target,
                        value=value,
                        timestamp=None,
                    )
                )

        return steps
