"""Tests for document extraction (T1: normalizer.extract)."""

from __future__ import annotations

import io
import subprocess
import sys
from pathlib import Path

import docx
import pytest
from PIL import Image
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from normalizer.extract import extract_document


def get_tiny_png() -> bytes:
    """Generate a 1x1 synthetic PNG in memory."""
    img = Image.new("RGB", (1, 1), color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def make_narrative_docx(
    file_path: Path,
    table_data: dict[str, str] | None = None,
    paragraphs: list[tuple[str, bytes | None] | str] | None = None,
) -> Path:
    """Create a synthetic .docx simulating a QA narrative specification."""
    doc = docx.Document()

    if table_data:
        table = doc.add_table(rows=2, cols=len(table_data))
        for col_idx, (k, v) in enumerate(table_data.items()):
            table.cell(0, col_idx).text = k
            table.cell(1, col_idx).text = v

    if paragraphs:
        for item in paragraphs:
            if isinstance(item, str):
                doc.add_paragraph(item)
            elif isinstance(item, tuple):
                text, img_bytes = item
                p = doc.add_paragraph(text if text else "")
                if img_bytes:
                    p.add_run().add_picture(io.BytesIO(img_bytes))

    doc.save(str(file_path))
    return file_path


def make_synthetic_pdf(file_path: Path, lines: list[str]) -> Path:
    """Create a synthetic PDF with text lines."""
    c = canvas.Canvas(str(file_path), pagesize=letter)
    y = 750
    for line in lines:
        c.drawString(100, y, line)
        y -= 20
    c.save()
    return file_path


@pytest.mark.smoke
def test_extract_docx_two_cases_line_ranges(tmp_path: Path) -> None:
    """Two cases in a docx produce two entries in cases with correct start_line and end_line."""
    doc_path = tmp_path / "two_cases.docx"
    out_dir = tmp_path / "out_two_cases"

    make_narrative_docx(
        doc_path,
        table_data={"Módulo": "Reportes", "Descripción corta": "Prueba de dos casos"},
        paragraphs=[
            "Link: https://example.com/app",
            "Caso 1: Primer caso de prueba",
            "Iniciamos sesión en el sistema",
            "Caso 2: Segundo caso de prueba",
            "Verificamos que la pantalla cargue",
        ],
    )

    dump = extract_document(doc_path, out_dir)

    assert len(dump["cases"]) == 2
    case_1 = dump["cases"][0]
    case_2 = dump["cases"][1]

    assert case_1["n"] == 1
    assert "Primer caso de prueba" in case_1["title"]
    assert case_1["start_line"] == 2
    assert case_1["end_line"] == 3

    assert case_2["n"] == 2
    assert "Segundo caso de prueba" in case_2["title"]
    assert case_2["start_line"] == 4
    assert case_2["end_line"] == 5

    assert (out_dir / "dump.json").exists()


def test_extract_docx_no_cases_fallback_to_table_title(tmp_path: Path) -> None:
    """Document without 'Caso N' creates a single case with table title."""
    doc_path = tmp_path / "no_cases.docx"
    out_dir = tmp_path / "out_no_cases"

    make_narrative_docx(
        doc_path,
        table_data={"Descripción completa": "Mantenimiento general de usuarios"},
        paragraphs=[
            "Link: https://example.com/app",
            "Iniciamos sesión en el sistema",
            "Pulsamos el botón Guardar",
        ],
    )

    dump = extract_document(doc_path, out_dir)

    assert len(dump["cases"]) == 1
    single_case = dump["cases"][0]
    assert single_case["n"] == 1
    assert single_case["title"] == "Mantenimiento general de usuarios"
    assert single_case["start_line"] == 1
    assert single_case["end_line"] == 3


@pytest.mark.smoke
def test_extract_docx_image_anchoring(tmp_path: Path) -> None:
    """An image in an image-only paragraph is anchored to the closest preceding text line."""
    doc_path = tmp_path / "with_image.docx"
    out_dir = tmp_path / "out_image"

    make_narrative_docx(
        doc_path,
        paragraphs=[
            "Caso 1: Flujo con evidencia",
            "Con la siguiente información:",
            ("", get_tiny_png()),
            "Verificamos el mensaje de éxito",
        ],
    )

    dump = extract_document(doc_path, out_dir)

    # Paragraphs should only contain the 3 text lines
    assert len(dump["paragraphs"]) == 3
    line_2 = dump["paragraphs"][1]
    assert line_2["text"] == "Con la siguiente información:"
    assert len(line_2["images"]) == 1

    img_rel_path = line_2["images"][0]
    saved_img = out_dir / img_rel_path
    assert saved_img.exists()
    assert saved_img.stat().st_size > 0

    # Line 1 and Line 3 should have no images
    assert dump["paragraphs"][0]["images"] == []
    assert dump["paragraphs"][2]["images"] == []


def test_extract_docx_negative_no_images(tmp_path: Path) -> None:
    """Document with only text produces empty images lists for all paragraphs."""
    doc_path = tmp_path / "no_images.docx"
    out_dir = tmp_path / "out_no_images"

    make_narrative_docx(
        doc_path,
        paragraphs=[
            "Caso 1: Sin imágenes",
            "Texto simple paso 1",
            "Texto simple paso 2",
        ],
    )

    dump = extract_document(doc_path, out_dir)

    for p in dump["paragraphs"]:
        assert p["images"] == []


def test_extract_cli_execution_success_and_negative(tmp_path: Path) -> None:
    """CLI execution via python -m normalizer.extract succeeds for docx and fails for invalid file."""
    doc_path = tmp_path / "cli_test.docx"
    out_dir = tmp_path / "cli_out"

    make_narrative_docx(
        doc_path,
        table_data={"Módulo": "CLI"},
        paragraphs=["Caso 1: CLI Test", "Link: https://example.com"],
    )

    cmd = [
        sys.executable,
        "-m",
        "normalizer.extract",
        str(doc_path),
        "--out",
        str(out_dir),
    ]
    res = subprocess.run(cmd, check=False, capture_output=True, text=True)
    assert res.returncode == 0
    assert (out_dir / "dump.json").exists()

    # Negative case: unsupported extension
    bad_file = tmp_path / "invalid.txt"
    bad_file.write_text("Not a docx or pdf", encoding="utf-8")
    bad_cmd = [
        sys.executable,
        "-m",
        "normalizer.extract",
        str(bad_file),
        "--out",
        str(tmp_path / "bad_out"),
    ]
    bad_res = subprocess.run(bad_cmd, check=False, capture_output=True, text=True)
    assert bad_res.returncode != 0


def test_extract_pdf_support_and_missing_file_negative(tmp_path: Path) -> None:
    """PDF extraction parses text and links; missing file raises error."""
    pdf_path = tmp_path / "sample.pdf"
    out_dir = tmp_path / "pdf_out"

    make_synthetic_pdf(
        pdf_path,
        [
            "Caso 1: PDF Extracted",
            "Link: https://example.com/pdf-portal",
            "Paso 1: Navegar a la página",
        ],
    )

    dump = extract_document(pdf_path, out_dir)

    assert len(dump["cases"]) >= 1
    assert "PDF Extracted" in dump["cases"][0]["title"]
    assert "https://example.com/pdf-portal" in dump["links"]

    # Negative: non-existent file
    with pytest.raises(FileNotFoundError):
        extract_document(tmp_path / "non_existent.pdf", out_dir)
