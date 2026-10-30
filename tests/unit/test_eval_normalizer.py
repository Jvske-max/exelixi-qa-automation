"""Unit tests for normalizer evaluation harness (T7)."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import docx
import pytest

from normalizer.schema import Case, NormalizedDoc, Step
from tools.eval_normalizer import evaluate_documents, main


@pytest.fixture
def synthetic_docs_and_normalized(tmp_path: Path) -> tuple[Path, Path, Path]:
    """Create a synthetic test suite with valid docx, corrupt file, and normalized JSON."""
    docs_dir = tmp_path / "docs"
    docs_dir.mkdir(parents=True, exist_ok=True)
    norm_dir = tmp_path / "normalized"
    norm_dir.mkdir(parents=True, exist_ok=True)
    out_file = tmp_path / "build" / "eval.md"

    # 1. doc_alpha.docx: Valid docx with sensitive strings and 2 cases
    doc_alpha_path = docs_dir / "doc_alpha.docx"
    doc_a = docx.Document()
    t_a = doc_a.add_table(rows=2, cols=4)
    t_a.rows[0].cells[0].text = "Fecha"
    t_a.rows[0].cells[1].text = "Módulo"
    t_a.rows[0].cells[2].text = "Identificador del caso"
    t_a.rows[0].cells[3].text = "Ambiente"
    t_a.rows[1].cells[0].text = "2026-10-28"
    t_a.rows[1].cells[1].text = "Seguridad"
    t_a.rows[1].cells[2].text = "R-ALPHA-01"
    t_a.rows[1].cells[3].text = "QA"

    doc_a.add_paragraph("Link: https://privado.internal/login")
    doc_a.add_paragraph("Caso 1: Creación de usuario confidencial Alfa")
    doc_a.add_paragraph("Objetivo de la Prueba: Registrar un usuario confidencial")
    doc_a.add_paragraph("Ingresamos al link https://privado.internal/login")
    doc_a.add_paragraph("Escribimos ${QA_USER_ADMIN} en el campo Usuario")
    doc_a.add_paragraph("Caso 2: Aprobación sensible")
    doc_a.add_paragraph("Pulsamos el botón Aprobar")
    doc_a.add_paragraph("Se solicita verificación de token OTP en pantalla")
    doc_a.add_paragraph("PALABRA_SECRETA_TOP_SECRET_123 y cédula V-99887766")
    doc_a.save(str(doc_alpha_path))

    # Normalized JSON for doc_alpha: R-ALPHA-01.json
    norm_doc = NormalizedDoc(
        meta={"identificador del caso": "R-ALPHA-01"},
        cases=[
            Case(
                n=1,
                title="Caso 1: Creación de usuario confidencial Alfa",
                objective="Registrar un usuario confidencial",
                url="https://privado.internal/login",
                credentials={"usuario": "${QA_USER_ADMIN}"},
                steps=[
                    Step(
                        action="navigate",
                        target="https://privado.internal/login",
                        source_quote="Ingresamos al link https://privado.internal/login",
                        source_ref="Caso 1, línea 4",
                    ),
                    Step(
                        action="fill",
                        target="Usuario",
                        value="${QA_USER_ADMIN}",
                        source_quote="Escribimos ${QA_USER_ADMIN} en el campo Usuario",
                        source_ref="Caso 1, línea 5",
                    ),
                ],
                gaps=[],
            ),
            Case(
                n=2,
                title="Caso 2: Aprobación sensible",
                objective="Aprobación sensible",
                url="https://privado.internal/login",
                credentials={},
                steps=[
                    Step(
                        action="click",
                        target="Aprobar",
                        source_quote="Pulsamos el botón Aprobar",
                        source_ref="Caso 2, línea 7",
                    ),
                    Step(
                        action="manual",
                        target="Se solicita verificación de token OTP en pantalla",
                        source_quote="Se solicita verificación de token OTP en pantalla",
                        source_ref="Caso 2, línea 8",
                    ),
                ],
                gaps=[],
            ),
        ],
    )
    (norm_dir / "R-ALPHA-01.json").write_text(norm_doc.model_dump_json(indent=2), encoding="utf-8")

    # 2. doc_corrupt.docx: Corrupt invalid binary file
    doc_corrupt_path = docs_dir / "doc_corrupt.docx"
    doc_corrupt_path.write_bytes(b"PK\x00\x00_NOT_A_VALID_DOCX_OR_ZIP_FILE_DATA_")

    # 3. doc_beta.docx: Valid docx without corresponding normalized JSON
    doc_beta_path = docs_dir / "doc_beta.docx"
    doc_b = docx.Document()
    t_b = doc_b.add_table(rows=2, cols=4)
    t_b.rows[0].cells[0].text = "Fecha"
    t_b.rows[0].cells[1].text = "Módulo"
    t_b.rows[0].cells[2].text = "Identificador del caso"
    t_b.rows[0].cells[3].text = "Ambiente"
    t_b.rows[1].cells[0].text = "2026-10-28"
    t_b.rows[1].cells[1].text = "Facturación"
    t_b.rows[1].cells[2].text = "R-BETA-02"
    t_b.rows[1].cells[3].text = "QA"

    doc_b.add_paragraph("Link: https://example.test/billing")
    doc_b.add_paragraph("Caso 1: Consulta de facturas")
    doc_b.add_paragraph("Iniciamos sesión con usuario ${QA_USER_ADMIN} y clave ${QA_PASS_ADMIN}")
    doc_b.add_paragraph("Paso 1: clic en Buscar Factura")
    doc_b.save(str(doc_beta_path))

    return docs_dir, norm_dir, out_file


@pytest.mark.smoke
def test_eval_normalizer_synthetic_suite(synthetic_docs_and_normalized: tuple[Path, Path, Path]) -> None:
    """Verify evaluation table format, strict privacy anti-leak, error resilience, and numeric metrics."""
    docs_dir, norm_dir, out_file = synthetic_docs_and_normalized

    # Positive run: evaluate synthetic dataset
    report_md = evaluate_documents(docs_dir=docs_dir, out_path=out_file, normalized_dir=norm_dir)

    assert out_file.exists(), "Output markdown report was not written"
    saved_content = out_file.read_text(encoding="utf-8")
    assert saved_content == report_md

    # (a) Check table row count: exactly one row per evaluated doc + header + separator
    lines = [ln.strip() for ln in saved_content.splitlines() if ln.strip().startswith("|")]
    assert len(lines) >= 3, "Table should have header, separator, and data rows"
    header_line = lines[0]
    separator_line = lines[1]
    data_rows = lines[2:]

    assert len(data_rows) == 3, f"Expected exactly 3 data rows for 3 docs, got {len(data_rows)}"
    assert "| #" in header_line or "| N" in header_line or "| #" in header_line
    assert "---" in separator_line

    # (b) Strict anti-leak / privacy verification: NO titles, text, filenames, or secrets in report
    forbidden_terms = [
        "doc_alpha",
        "doc_corrupt",
        "doc_beta",
        "Creación de usuario confidencial Alfa",
        "Objetivo de la Prueba",
        "https://privado.internal/login",
        "https://example.test/billing",
        "QA_ADMIN_USER",
        "PALABRA_SECRETA_TOP_SECRET_123",
        "V-99887766",
        "Aprobación sensible",
        "Consulta de facturas",
    ]
    for term in forbidden_terms:
        assert term not in saved_content, f"Privacy violation: '{term}' leaked in eval report!"

    # (c) Corrupt file generates an ERROR row without halting evaluation of remaining docs
    # Alphabetical order: doc_alpha (1), doc_beta (2), doc_corrupt (3) or similar
    # Find row with ERROR
    error_rows = [r for r in data_rows if "ERROR" in r]
    assert len(error_rows) == 1, "Expected exactly 1 error row for corrupt document"
    error_row = error_rows[0]
    error_cells = [c.strip() for c in error_row.split("|")[1:-1]]
    # Ensure error row has ERROR in status and doesn't abort other rows
    assert "ERROR" in error_cells

    # (d) Numeric metrics verification for doc_alpha (with normalized JSON)
    # doc_alpha has 2 cases:
    # Case 1: 2 exec steps, 0 manual. Quotes match original text.
    # Case 2: 1 exec step, 1 manual step. Quotes match original text.
    # Total: 2 cases, 3 exec steps, 1 manual step, 100.0% valid quotes, READY: 2
    alpha_row = data_rows[0]  # doc_alpha is first alphabetically
    alpha_cells = [c.strip() for c in alpha_row.split("|")[1:-1]]

    assert alpha_cells[0] == "1", "Order number should be 1"
    assert alpha_cells[1] == "2", "Total cases should be 2"
    assert alpha_cells[2] == "3", "Executable steps should be 3"
    assert alpha_cells[3] == "1", "Manual steps should be 1"
    assert alpha_cells[5] == "100.0%", "Quote percentage should be 100.0%"
    assert "READY" in alpha_cells[6], "Status should indicate READY"


@pytest.mark.smoke
def test_eval_normalizer_fabricated_quote_negative_metric(tmp_path: Path) -> None:
    """Negative grounding test: fabricated source_quote calculates < 100% valid quotes."""
    docs_dir = tmp_path / "docs"
    docs_dir.mkdir(parents=True, exist_ok=True)
    norm_dir = tmp_path / "normalized"
    norm_dir.mkdir(parents=True, exist_ok=True)
    out_file = tmp_path / "eval.md"

    # Create docx with specific text
    doc_path = docs_dir / "sample_quote.docx"
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=4)
    t.rows[0].cells[2].text = "Identificador del caso"
    t.rows[1].cells[2].text = "R-QUOTE-01"
    doc.add_paragraph("Link: https://example.test/app")
    doc.add_paragraph("Caso 1: Verificación de citas")
    doc.add_paragraph("Paso 1: Abrir navegador")
    doc.save(str(doc_path))

    # Normalized doc has 2 steps: 1 valid quote, 1 fabricated quote
    norm_doc = NormalizedDoc(
        meta={"identificador del caso": "R-QUOTE-01"},
        cases=[
            Case(
                n=1,
                title="Caso 1",
                objective="Test quotes",
                url="https://example.test/app",
                steps=[
                    Step(
                        action="navigate",
                        target="https://example.test/app",
                        source_quote="Paso 1: Abrir navegador",  # Valid (verbatim)
                        source_ref="Caso 1, línea 3",
                    ),
                    Step(
                        action="click",
                        target="Boton Inexistente",
                        source_quote="Esta frase NUNCA estuvo en el documento original",  # Fabricated!
                        source_ref="Caso 1, línea 4",
                    ),
                ],
            )
        ],
    )
    (norm_dir / "sample_quote.json").write_text(norm_doc.model_dump_json(indent=2), encoding="utf-8")

    report_md = evaluate_documents(docs_dir=docs_dir, out_path=out_file, normalized_dir=norm_dir)

    lines = [ln.strip() for ln in report_md.splitlines() if ln.strip().startswith("|")]
    data_rows = lines[2:]
    assert len(data_rows) == 1
    cells = [c.strip() for c in data_rows[0].split("|")[1:-1]]

    # 1 valid quote out of 2 steps = 50.0%
    assert cells[5] == "50.0%", f"Expected 50.0% for 1/2 valid quotes, got {cells[5]}"


@pytest.mark.smoke
def test_eval_normalizer_cli_invocation(synthetic_docs_and_normalized: tuple[Path, Path, Path]) -> None:
    """Verify CLI execution via python -m tools.eval_normalizer."""
    docs_dir, norm_dir, out_file = synthetic_docs_and_normalized

    cmd = [
        sys.executable,
        "-m",
        "tools.eval_normalizer",
        "--docs",
        str(docs_dir),
        "--out",
        str(out_file),
        "--normalized-dir",
        str(norm_dir),
    ]

    res = subprocess.run(cmd, check=False, capture_output=True, text=True)
    assert res.returncode == 0, f"CLI exited with error: {res.stderr}"
    assert out_file.exists()
    assert "Evaluación de Normalización" in out_file.read_text(encoding="utf-8")


@pytest.mark.smoke
def test_eval_normalizer_missing_docs_dir_negative(tmp_path: Path) -> None:
    """Negative test: Non-existent --docs directory raises FileNotFoundError or returns code 1."""
    non_existent = tmp_path / "does_not_exist"
    out_file = tmp_path / "eval.md"

    with pytest.raises(FileNotFoundError):
        evaluate_documents(docs_dir=non_existent, out_path=out_file)

    ret = main(["--docs", str(non_existent), "--out", str(out_file)])
    assert ret != 0


@pytest.mark.smoke
def test_eval_normalizer_with_pdf_and_corrupt_pdf(tmp_path: Path) -> None:
    """Verify evaluation of valid PDF and corrupt PDF document handling."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    docs_dir = tmp_path / "pdf_docs"
    docs_dir.mkdir(parents=True, exist_ok=True)
    out_file = tmp_path / "eval_pdf.md"

    # Valid PDF
    valid_pdf = docs_dir / "doc_valid.pdf"
    c = canvas.Canvas(str(valid_pdf), pagesize=letter)
    c.drawString(100, 750, "Link: https://example.test/pdf")
    c.drawString(100, 725, "Caso 1: Verificación de PDF")
    c.drawString(100, 700, "Paso 1: Abrir pantalla principal")
    c.save()

    # Corrupt PDF
    corrupt_pdf = docs_dir / "doc_corrupt.pdf"
    corrupt_pdf.write_bytes(b"%PDF-1.4\n%corrupt_byte_sequence_that_fails_parsing_completely")

    report_md = evaluate_documents(docs_dir=docs_dir, out_path=out_file, normalized_dir=tmp_path)

    lines = [ln.strip() for ln in report_md.splitlines() if ln.strip().startswith("|")]
    data_rows = lines[2:]
    assert len(data_rows) == 2, f"Expected 2 rows for 2 pdfs, got {len(data_rows)}"

    # Confirm corrupt PDF row contains ERROR and valid PDF row did not fail
    error_rows = [r for r in data_rows if "ERROR" in r]
    assert len(error_rows) == 1, "Expected 1 error row for corrupt PDF"

    # Strict anti-leak: verify no terms leak
    assert "doc_valid" not in report_md
    assert "doc_corrupt" not in report_md
    assert "Verificación de PDF" not in report_md


@pytest.mark.smoke
def test_eval_normalizer_empty_directory(tmp_path: Path) -> None:
    """Verify evaluation on an empty directory produces header and 0 data rows."""
    empty_docs = tmp_path / "empty_docs"
    empty_docs.mkdir(parents=True, exist_ok=True)
    out_file = tmp_path / "eval_empty.md"

    report_md = evaluate_documents(docs_dir=empty_docs, out_path=out_file)

    lines = [ln.strip() for ln in report_md.splitlines() if ln.strip().startswith("|")]
    assert len(lines) == 2, f"Expected 2 lines (header + separator), got {len(lines)}"

