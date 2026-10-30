"""Unit tests for normalizer.render and CLI render subcommand (T4)."""

from __future__ import annotations

import re
from pathlib import Path

import docx
import pytest

from normalizer.cli import main as cli_main
from normalizer.gaps import Gap
from normalizer.render import _format_credential, render_case
from normalizer.schema import Case, NormalizedDoc, Step
from parsers.docx_parser import DocxParser


@pytest.fixture
def synthetic_normalized_doc() -> NormalizedDoc:
    """Fixture providing a synthetic NormalizedDoc with valid canonical steps."""
    return NormalizedDoc(
        meta={
            "identificador del caso": "SYNTH-001",
            "módulo": "Facturación",
            "ambiente de prueba": "QA",
            "analista programador": "QA Automation",
            "analista de prueba": "QA Automation",
            "descripción corta": "Caso sintético completo",
        },
        cases=[
            Case(
                n=1,
                title="Flujo completo de facturación",
                objective="Verificar emisión de comprobante sintético",
                url="https://qa.example.test/facturacion",
                credentials={
                    "usuario": "${QA_USER_ADMIN}",
                    "clave": "${QA_PASS_ADMIN}",
                },
                steps=[
                    Step(
                        action="navigate",
                        target="https://qa.example.test/facturacion",
                        source_quote="Ingresamos al portal",
                        source_ref="Párrafo 1",
                    ),
                    Step(
                        action="manual",
                        target="Adjuntar certificado fiscal digital",
                        source_quote="Se adjunta el certificado digital en disco",
                        source_ref="Párrafo 2",
                    ),
                    Step(
                        action="fill",
                        target="Campo Código",
                        value="FAC-2026",
                        source_quote="Colocamos el código FAC-2026",
                        source_ref="Párrafo 3",
                    ),
                    Step(
                        action="click",
                        target="Boton Emitir",
                        source_quote="Pulsamos el botón Emitir",
                        source_ref="Párrafo 4",
                    ),
                    Step(
                        action="verify",
                        target="Comprobante generado exitosamente",
                        source_quote="Se muestra mensaje Comprobante generado exitosamente",
                        source_ref="Párrafo 5",
                    ),
                ],
                gaps=[],
            )
        ],
    )


@pytest.mark.smoke
def test_render_roundtrip_with_docx_parser(tmp_path: Path, synthetic_normalized_doc: NormalizedDoc) -> None:
    """Roundtrip test: render_case generates a .docx that DocxParser parses correctly."""
    case = synthetic_normalized_doc.cases[0]

    # Positive case: valid docx rendered and parsed back
    docx_path = render_case(synthetic_normalized_doc, case, tmp_path)
    assert docx_path.exists()
    assert docx_path.suffix == ".docx"

    parsed = DocxParser.parse(docx_path)

    # 1. Check executable step count matches exactly (4 executable, 1 manual ignored)
    executable_steps = [s for s in case.steps if s.action != "manual"]
    assert len(parsed.steps) == len(executable_steps)
    assert len(parsed.steps) == 4

    # 2. Check actions, targets, and values match
    assert parsed.steps[0].action == "navigate"
    assert parsed.steps[0].target == "https://qa.example.test/facturacion"
    assert parsed.steps[0].value is None

    assert parsed.steps[1].action == "fill"
    assert parsed.steps[1].target == "Campo Código"
    assert parsed.steps[1].value == "FAC-2026"

    assert parsed.steps[2].action == "click"
    assert parsed.steps[2].target == "Boton Emitir"
    assert parsed.steps[2].value is None

    assert parsed.steps[3].action == "verify"
    assert parsed.steps[3].target == "Comprobante generado exitosamente"
    assert parsed.steps[3].value is None

    # 3. Check manual steps are NOT in parsed.steps
    targets_and_texts = [f"{s.target} {s.raw_text}" for s in parsed.steps]
    for text in targets_and_texts:
        assert "certificado fiscal" not in text.lower()
        assert "pendiente manual" not in text.lower()

    # Negative case: case with only manual steps produces 0 executable steps in DocxParser
    manual_case = Case(
        n=2,
        title="Solo flujo manual",
        objective="Probar pasarela externa",
        steps=[
            Step(
                action="manual",
                target="Confirmar pago en banco",
                source_quote="Se valida pago en el banco",
                source_ref="Párrafo 1",
            )
        ],
    )
    manual_docx = render_case(synthetic_normalized_doc, manual_case, tmp_path)
    parsed_manual = DocxParser.parse(manual_docx)
    assert len(parsed_manual.steps) == 0


@pytest.mark.smoke
def test_render_table_and_status(tmp_path: Path, synthetic_normalized_doc: NormalizedDoc) -> None:
    """Table 'Identificación' contains canonical headers, metadata, and normalization status."""
    case_ready = synthetic_normalized_doc.cases[0]

    # Positive case: READY status in table
    docx_ready = render_case(synthetic_normalized_doc, case_ready, tmp_path)
    doc_docx = docx.Document(str(docx_ready))
    assert len(doc_docx.tables) >= 1

    table = doc_docx.tables[0]
    assert len(table.rows) >= 2

    headers = [c.text.strip().lower() for c in table.rows[0].cells]
    values = [c.text.strip() for c in table.rows[1].cells]

    header_to_val = dict(zip(headers, values))
    assert "estado de normalización" in header_to_val
    assert header_to_val["estado de normalización"] == "READY"
    assert header_to_val["módulo"] == "Facturación"
    assert header_to_val["identificador del caso"] == "SYNTH-001"

    # Negative/Edge cases: NEEDS_INFO and UNSUPPORTED status in table
    case_needs_info = Case(
        n=3,
        title="Caso incompleto",
        objective="Falta URL",
        steps=[
            Step(
                action="click",
                target="Boton",
                source_quote="clic",
                source_ref="P1",
            )
        ],
        gaps=[Gap(type="missing_url", severity="blocking", where="Caso 3", ask="Añadir URL")],
    )
    docx_needs_info = render_case(synthetic_normalized_doc, case_needs_info, tmp_path)
    table_needs_info = docx.Document(str(docx_needs_info)).tables[0]
    headers_ni = [c.text.strip().lower() for c in table_needs_info.rows[0].cells]
    vals_ni = [c.text.strip() for c in table_needs_info.rows[1].cells]
    dict_ni = dict(zip(headers_ni, vals_ni))
    assert dict_ni["estado de normalización"] == "NEEDS_INFO"

    case_unsupported = Case(
        n=4,
        title="Caso no soportado",
        objective="Solo OTP",
        steps=[
            Step(
                action="manual",
                target="OTP",
                source_quote="otp",
                source_ref="P1",
            )
        ],
    )
    docx_unsupported = render_case(synthetic_normalized_doc, case_unsupported, tmp_path)
    table_unsupported = docx.Document(str(docx_unsupported)).tables[0]
    headers_uns = [c.text.strip().lower() for c in table_unsupported.rows[0].cells]
    vals_uns = [c.text.strip() for c in table_unsupported.rows[1].cells]
    dict_uns = dict(zip(headers_uns, vals_uns))
    assert dict_uns["estado de normalización"] == "UNSUPPORTED"


@pytest.mark.smoke
def test_render_gaps_report(tmp_path: Path, synthetic_normalized_doc: NormalizedDoc) -> None:
    """NEEDS_INFO case generates <stem>.gaps.md ordered by severity (blocking -> warning -> info)."""
    # Positive case: gaps with unordered severities
    gaps_case = Case(
        n=5,
        title="Caso con faltantes",
        objective="Comprobar reporte gaps",
        steps=[
            Step(
                action="click",
                target="Boton",
                source_quote="clic",
                source_ref="P1",
            )
        ],
        gaps=[
            Gap(
                type="unsupported_flow",
                severity="info",
                where="Caso 5, línea 10",
                ask="Paso manual sugerido",
            ),
            Gap(
                type="missing_url",
                severity="blocking",
                where="Caso 5",
                ask="Añadir Link: https://qa.test",
            ),
            Gap(
                type="missing_expected",
                severity="warning",
                where="Caso 5, línea 4",
                ask="Escribir texto esperado",
            ),
        ],
    )

    docx_path = render_case(synthetic_normalized_doc, gaps_case, tmp_path)
    gaps_path = docx_path.with_name(f"{docx_path.stem}.gaps.md")

    assert gaps_path.exists()
    content = gaps_path.read_text(encoding="utf-8")

    # Verify severity ordering: blocking appears before warning, warning before info
    idx_blocking = content.find("blocking")
    idx_warning = content.find("warning")
    idx_info = content.find("info")

    assert idx_blocking != -1
    assert idx_warning != -1
    assert idx_info != -1
    assert idx_blocking < idx_warning < idx_info

    # Verify attributes are present
    assert "missing_url" in content
    assert "Añadir Link: https://qa.test" in content
    assert "Caso 5, línea 10" in content

    # Negative case: case with READY status and no gaps does NOT generate .gaps.md
    ready_case = synthetic_normalized_doc.cases[0]
    docx_ready = render_case(synthetic_normalized_doc, ready_case, tmp_path)
    gaps_ready = docx_ready.with_name(f"{docx_ready.stem}.gaps.md")
    assert not gaps_ready.exists()


@pytest.mark.smoke
def test_render_case_paragraph_order(tmp_path: Path, synthetic_normalized_doc: NormalizedDoc) -> None:
    """Paragraph 'Caso N: <title>' must be the very first paragraph in the document containing 'caso'."""
    case = synthetic_normalized_doc.cases[0]
    docx_path = render_case(synthetic_normalized_doc, case, tmp_path)

    doc = docx.Document(str(docx_path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Find all paragraphs containing 'caso' (case-insensitive)
    caso_paragraphs = [p for p in paragraphs if re.search(r"\bcaso\b", p, re.IGNORECASE)]

    # Positive check: at least one exists and the first is 'Caso <n>: <title>'
    assert len(caso_paragraphs) > 0
    assert caso_paragraphs[0].startswith(f"Caso {case.n}:")

    # Negative check: neither the table header nor link should appear as a preceding 'caso' paragraph
    idx_first_caso = paragraphs.index(caso_paragraphs[0])
    for prev_p in paragraphs[:idx_first_caso]:
        assert not re.search(r"\bcaso\b", prev_p, re.IGNORECASE)


@pytest.mark.smoke
def test_render_safe_filenames_and_cli(tmp_path: Path, synthetic_normalized_doc: NormalizedDoc) -> None:
    """CLI subcommand render generates safe files following <caso_id>-caso<N>.docx."""
    out_dir = tmp_path / "output_normalized"
    json_path = tmp_path / "synthetic_doc.json"

    json_path.write_text(synthetic_normalized_doc.model_dump_json(indent=2), encoding="utf-8")

    # Positive CLI invocation
    ret_code = cli_main(["render", str(json_path), "--out", str(out_dir)])
    assert ret_code == 0

    expected_docx = out_dir / "SYNTH-001-caso1.docx"
    assert expected_docx.exists()

    # Check filename does not contain personal data, RIFs, or credentials
    filename = expected_docx.name
    assert not re.search(r"(?:[VvEeJjGg]-\d+|\bpass\b|\bclave\b)", filename, re.IGNORECASE)

    # Negative case 1: case_id with unsafe characters/path traversal is sanitized
    doc_unsafe = NormalizedDoc(
        meta={"identificador del caso": "../../MALICIOUS:ID/*#"},
        cases=[
            Case(
                n=7,
                title="Prueba de sanitización",
                objective="Verificar nombre de archivo seguro",
                steps=[
                    Step(
                        action="click",
                        target="Boton",
                        source_quote="clic",
                        source_ref="P1",
                    )
                ],
            )
        ],
    )
    docx_sanitized = render_case(doc_unsafe, doc_unsafe.cases[0], out_dir)
    assert docx_sanitized.exists()
    assert ".." not in docx_sanitized.name
    assert ":" not in docx_sanitized.name
    assert "*" not in docx_sanitized.name
    assert docx_sanitized.name.endswith("-caso7.docx")

    # Negative case 2: CLI invocation with non-existent file returns non-zero code
    ret_bad = cli_main(["render", str(tmp_path / "non_existent.json"), "--out", str(out_dir)])
    assert ret_bad != 0


@pytest.mark.smoke
def test_render_never_leaks_raw_credential_value(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    """Verify that literal password 'zq9-clave-prueba' is never leaked into .docx, .gaps.md, filenames, or CLI output."""
    raw_secret = "zq9-clave-prueba"
    doc = NormalizedDoc(
        meta={"identificador del caso": "R-LEAK-TEST"},
        cases=[
            Case(
                n=1,
                title="Prueba de fuga de clave",
                objective="Verificar que el render nunca filtra claves literales",
                credentials={"usuario": "operador_user", "clave": raw_secret},
                steps=[
                    Step(
                        action="fill",
                        target="Clave",
                        value=raw_secret,
                        source_quote="Escribimos la clave secreta",
                        source_ref="Línea 2",
                    ),
                ],
                gaps=[
                    Gap(
                        type="credentials_in_plaintext",
                        severity="warning",
                        where="Caso 1, línea 2",
                        ask="Mover a variable de entorno",
                    )
                ],
                status="NEEDS_INFO",
            )
        ],
    )
    out_dir = tmp_path / "leak_check_out"
    json_path = tmp_path / "doc_with_secret.json"
    json_path.write_text(doc.model_dump_json(indent=2), encoding="utf-8")

    code = cli_main(["render", str(json_path), "--out", str(out_dir)])
    assert code == 0

    captured = capsys.readouterr()
    cli_output = captured.out + captured.err
    assert raw_secret.lower() not in cli_output.lower()

    # Check generated files
    rendered_files = list(out_dir.iterdir())
    assert len(rendered_files) >= 1

    for f in rendered_files:
        # Check filename
        assert raw_secret.lower() not in f.name.lower()
        if f.suffix == ".docx":
            d = docx.Document(str(f))
            full_text = " ".join(p.text for p in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    full_text += " " + " ".join(c.text for c in row.cells)
            assert raw_secret.lower() not in full_text.lower()
            # Ensure env var is derived from default / role, not raw password
            assert "QA_PASS_DEFAULT" in full_text
            assert "QA_USER_DEFAULT" in full_text
        elif f.suffix == ".md":
            md_text = f.read_text(encoding="utf-8")
            assert raw_secret.lower() not in md_text.lower()


@pytest.mark.smoke
def test_render_sanitizes_analyst_names(tmp_path: Path) -> None:
    """Verify that analyst fields in rendered docx are always sanitized to generic 'Exelixi QA'."""
    # Negative test case: doc.meta contains real personal names (Juan Pérez, María Gómez)
    personal_names = ["Juan Pérez", "María Gómez"]
    doc_with_personal_names = NormalizedDoc(
        meta={
            "identificador del caso": "SEC-ANALYST-01",
            "módulo": "Seguridad",
            "ambiente": "QA",
            "analista programador": personal_names[0],
            "analista de prueba": personal_names[1],
            "analista": "Carlos Ruiz",
        },
        cases=[
            Case(
                n=1,
                title="Prueba de sanitización de analistas",
                objective="Verificar que nunca se filtren nombres de personas",
                steps=[
                    Step(
                        action="click",
                        target="Boton Continuar",
                        source_quote="clic en Continuar",
                        source_ref="P1",
                    )
                ],
            )
        ],
    )

    docx_path = render_case(doc_with_personal_names, doc_with_personal_names.cases[0], tmp_path)
    assert docx_path.exists()

    d = docx.Document(str(docx_path))
    table = d.tables[0]
    headers = [c.text.strip().lower() for c in table.rows[0].cells]
    values = [c.text.strip() for c in table.rows[1].cells]
    table_dict = dict(zip(headers, values))

    # Assert that personal names are NEVER in the table or document
    for name in personal_names + ["Carlos Ruiz"]:
        assert name.lower() not in table_dict["analista programador"].lower()
        assert name.lower() not in table_dict["analista de prueba"].lower()
        assert name.lower() not in " ".join(values).lower()

    # Positive test: Assert institutional generic name is used instead
    assert table_dict["analista programador"] == "Exelixi QA"
    assert table_dict["analista de prueba"] == "Exelixi QA"


@pytest.mark.smoke
def test_render_masks_known_passwords_in_generic_fill_target(tmp_path: Path) -> None:
    """Verify that known passwords are masked in generic fill steps even if target is generic (e.g. campo_secreto)."""
    doc = NormalizedDoc(
        meta={"identificador del caso": "SEC-GENERIC-01", "módulo": "Seguridad"},
        cases=[
            Case(
                n=1,
                title="Prueba de enmascaramiento con target genérico",
                objective="Verificar que claves detectadas no se estampen en docx",
                steps=[
                    Step(
                        action="fill",
                        target="campo_secreto",
                        value="clave_detectada_123",
                        source_quote="Ingresamos clave_detectada_123",
                        source_ref="P1",
                    ),
                    Step(
                        action="fill",
                        target="campo_codigo",
                        value="CODIGO_VALIDO_777",
                        source_quote="Ingresamos código 777",
                        source_ref="P2",
                    ),
                ],
                gaps=[
                    Gap(
                        type="credentials_in_plaintext",
                        severity="warning",
                        where="Caso 1, línea 1",
                        ask="Mover clave a .env",
                    )
                ],
            )
        ],
    )
    case = doc.cases[0]

    # Negative check: secret must be masked with ${QA_PASS_DEFAULT} and literal secret MUST NOT appear
    docx_path = render_case(
        doc,
        case,
        tmp_path,
        known_passwords={"clave_detectada_123"},
    )
    assert docx_path.exists()

    d = docx.Document(str(docx_path))
    doc_text = " ".join(p.text for p in d.paragraphs)

    assert "clave_detectada_123" not in doc_text
    assert "${QA_PASS_DEFAULT}" in doc_text

    # Positive check: normal non-secret fill value remains untouched
    assert "CODIGO_VALIDO_777" in doc_text


@pytest.mark.smoke
def test_format_credential_rejects_lowercase_env_vars() -> None:
    """Verify that _format_credential rejects lowercase env vars and falls back to uppercase standard format."""
    # Negative check 1: lowercase env var without role must not be returned literally; returns ${QA_USER_DEFAULT}
    res_default = _format_credential("${qa_user_test}", "QA_USER")
    assert res_default != "${qa_user_test}"
    assert res_default == "${QA_USER_DEFAULT}"

    # Negative check 2: lowercase env var with role returns uppercase role-based variable
    res_role = _format_credential("${qa_user_test}", "QA_USER", role="admin")
    assert res_role != "${qa_user_test}"
    assert res_role == "${QA_USER_ADMIN}"

    # Positive check 1: valid uppercase env var is preserved
    assert _format_credential("${QA_USER_SUPERVISOR}", "QA_USER") == "${QA_USER_SUPERVISOR}"

    # Positive check 2: None value derives default uppercase variable
    assert _format_credential(None, "QA_PASS") == "${QA_PASS_DEFAULT}"


@pytest.mark.smoke
def test_render_manual_step_sanitizes_pii_in_source_quote_fallback(tmp_path: Path) -> None:
    """Verify that manual steps falling back to source_quote sanitize PII (cédula, email, etc.)."""
    doc = NormalizedDoc(
        meta={"identificador del caso": "PII-MANUAL-01", "módulo": "Seguridad"},
        cases=[
            Case(
                n=1,
                title="Caso con PII en cita manual",
                objective="Verificar sanitización de PII en fallback de cita de paso manual",
                url="https://qa.example.test/login",
                credentials={"usuario": "${QA_USER_DEFAULT}", "clave": "${QA_PASS_DEFAULT}"},
                steps=[
                    Step(
                        action="manual",
                        target="",
                        value="",
                        source_quote="Ingresamos la cédula V-12345678 y correo usuario_qa@empresa.com",
                        source_ref="Párrafo 1",
                    )
                ],
            )
        ],
    )
    case = doc.cases[0]
    docx_path = render_case(doc, case, tmp_path)
    assert docx_path.exists()

    d = docx.Document(str(docx_path))
    manual_paragraphs = [p.text for p in d.paragraphs if p.text.startswith("Pendiente manual 1:")]
    assert len(manual_paragraphs) == 1
    p_text = manual_paragraphs[0]

    # Negative check 1: literal PII values must not appear in rendered document
    assert "V-12345678" not in p_text
    assert "usuario_qa@empresa.com" not in p_text

    # Positive check 1: protected tags replace sensitive data
    assert any(tag in p_text for tag in ("[DATO_PROTEGIDO]", "[CEDULA_PROTEGIDA]"))
    assert any(tag in p_text for tag in ("[DATO_PROTEGIDO]", "[EMAIL_PROTEGIDO]"))

    # Positive check 2: non-sensitive narrative text is preserved
    assert "Ingresamos la cédula" in p_text
    assert "correo" in p_text
    assert "Pendiente manual 1:" in p_text


@pytest.mark.smoke
def test_render_gaps_report_sanitizes_pii(tmp_path: Path, synthetic_normalized_doc: NormalizedDoc) -> None:
    """Verify that _render_gaps_report sanitizes PII (cédula, email) in gaps.md output."""
    # 1. Negative case: gap description and where contain cédula and email
    gap_case_pii = Case(
        n=10,
        title="Caso con PII en faltantes",
        objective="Verificar sanitización de gaps.md",
        steps=[
            Step(
                action="click",
                target="Boton",
                source_quote="clic en boton",
                source_ref="Párrafo 1",
            )
        ],
        gaps=[
            Gap(
                type="missing_credentials",
                severity="blocking",
                where="Ubicación analista V-99887766",
                ask="Solicitar acceso a contacto@empresa.com para continuar",
            )
        ],
    )

    docx_path = render_case(synthetic_normalized_doc, gap_case_pii, tmp_path)
    gaps_path = docx_path.with_name(f"{docx_path.stem}.gaps.md")
    assert gaps_path.exists()
    content = gaps_path.read_text(encoding="utf-8")

    # Assert that literal PII is not present in .gaps.md
    assert "V-99887766" not in content, "Cédula should not be present in .gaps.md"
    assert "contacto@empresa.com" not in content, "Email should not be present in .gaps.md"
    # Assert that PII was sanitized with protected tags
    assert "[CEDULA_PROTEGIDA]" in content or "[DATO_PROTEGIDO]" in content
    assert "[EMAIL_PROTEGIDO]" in content or "[DATO_PROTEGIDO]" in content

    # 2. Positive case: normal gap without PII preserves narrative text
    gap_case_clean = Case(
        n=11,
        title="Caso limpio",
        objective="Verificar gaps limpios",
        steps=[
            Step(
                action="click",
                target="Boton",
                source_quote="clic en boton",
                source_ref="Párrafo 1",
            )
        ],
        gaps=[
            Gap(
                type="missing_url",
                severity="blocking",
                where="Pantalla Principal",
                ask="Indicar la URL de inicio del portal",
            )
        ],
    )
    docx_clean = render_case(synthetic_normalized_doc, gap_case_clean, tmp_path)
    gaps_clean = docx_clean.with_name(f"{docx_clean.stem}.gaps.md")
    assert gaps_clean.exists()
    clean_content = gaps_clean.read_text(encoding="utf-8")
    assert "Pantalla Principal" in clean_content
    assert "Indicar la URL de inicio del portal" in clean_content


@pytest.mark.smoke
def test_render_masks_quoted_known_password_in_generic_fill_step(
    tmp_path: Path, synthetic_normalized_doc: NormalizedDoc
) -> None:
    """Verify that quoted known password in generic fill step is masked with ${QA_PASS_DEFAULT}."""
    # Negative/masked case: quoted known password on generic target must be masked
    case_quoted_pass = Case(
        n=1,
        title="Quoted Password Masking Case",
        objective="Verificar enmascaramiento de clave entrecomillada en fill genérico",
        url="https://qa.example.test/portal",
        steps=[
            Step(
                action="fill",
                target="campo_codigo",
                value='"secreto_render_888"',
                source_quote='escribir "secreto_render_888" en campo_codigo',
                source_ref="Párrafo 1",
            )
        ],
    )
    docx_path = render_case(
        synthetic_normalized_doc,
        case_quoted_pass,
        tmp_path,
        known_passwords={"secreto_render_888"},
    )
    doc = docx.Document(str(docx_path))
    doc_text = "\n".join(p.text for p in doc.paragraphs)

    assert "${QA_PASS_DEFAULT}" in doc_text
    assert "secreto_render_888" not in doc_text

    # Positive/unmasked case: normal non-secret quoted value is not replaced by pass_default
    case_normal_val = Case(
        n=2,
        title="Normal Value Case",
        objective="Verificar que valor normal no sea enmascarado como clave",
        url="https://qa.example.test/portal",
        steps=[
            Step(
                action="fill",
                target="campo_codigo",
                value='"codigo_valido_123"',
                source_quote='escribir "codigo_valido_123" en campo_codigo',
                source_ref="Párrafo 2",
            )
        ],
    )
    docx_path_pos = render_case(
        synthetic_normalized_doc,
        case_normal_val,
        tmp_path,
        known_passwords={"secreto_render_888"},
    )
    doc_pos = docx.Document(str(docx_path_pos))
    pos_text = "\n".join(p.text for p in doc_pos.paragraphs)

    assert "codigo_valido_123" in pos_text
    assert "${QA_PASS_DEFAULT}" not in pos_text





