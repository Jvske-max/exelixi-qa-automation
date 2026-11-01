"""Unit tests for document ingestion parsers (.docx, .pdf, unified DocParser)."""

from pathlib import Path

import docx
import pytest
from reportlab.pdfgen import canvas

from parsers import DocParser, DocxParser, PdfParser
from parsers.models import StepDefinition, TestCase

PROD_DOCX_PATH = Path("docs-qa/sample_case.docx")


@pytest.fixture
def sample_docx_path(tmp_path: Path) -> Path:
    """Creates a lightweight sample .docx document with metadata and steps."""
    docx_file = tmp_path / "sample_case.docx"
    doc = docx.Document()
    
    # Metadata table
    table = doc.add_table(rows=5, cols=3)
    table.rows[0].cells[0].text = "Fecha"
    table.rows[0].cells[1].text = "Módulo"
    table.rows[0].cells[2].text = "Identificador del caso"
    
    table.rows[1].cells[0].text = "01/09/2026"
    table.rows[1].cells[1].text = "Pasarela"
    table.rows[1].cells[2].text = "R-055429"
    
    table.rows[2].cells[0].text = "Descripción corta"
    table.rows[2].cells[1].text = "Migración Pasarela"
    table.rows[2].cells[2].text = "Migración Pasarela"
    
    table.rows[3].cells[0].text = "Ambiente de prueba"
    table.rows[3].cells[1].text = "QA"
    table.rows[3].cells[2].text = "QA"
    
    table.rows[4].cells[0].text = "Analista Programador"
    table.rows[4].cells[1].text = "QA Tester"
    table.rows[4].cells[2].text = "QA Tester"
    
    # Paragraphs for URL, title, objective, and credentials
    doc.add_paragraph("Link: https://qa.example.test/login")
    doc.add_paragraph("Caso # 3: Usuario para Reportería")
    doc.add_paragraph("Objetivo de la Prueba")
    doc.add_paragraph("Verificar el funcionamiento de la creación y edición de usuarios.")
    doc.add_paragraph("Usuario: test_user")
    doc.add_paragraph("Clave: test_password")
    
    # Noise paragraphs
    doc.add_paragraph("Step 1 screenshot.")
    doc.add_paragraph("Previous Next")
    
    # Steps
    doc.add_paragraph(
        'Step 2: (9/2/2026 3:21:29 PM) User keyboard input on '
        '"Mundial de Seguros - Avast Secure Browser (window)" '
        'in "Mundial de Seguros - Avast Secure Browser" [... CTRL-V TABULACION ...]'
    )
    doc.add_paragraph("Step 2 screenshot.")
    doc.add_paragraph("Previous Next")
    doc.add_paragraph(
        'Step 3: (9/2/2026 3:21:33 PM) User left click on '
        '"Mantenimiento (link)" in "Mundial de Seguros - Avast Secure Browser"'
    )
    doc.add_paragraph(
        'Step 4: (9/2/2026 3:21:34 PM) User left click on '
        '"Seguridad (link)" in "Mundial de Seguros - Avast Secure Browser"'
    )
    
    doc.save(str(docx_file))
    return docx_file


@pytest.fixture
def sample_pdf_path(tmp_path: Path) -> Path:
    """Creates a sample .pdf document with metadata and steps."""
    pdf_file = tmp_path / "sample_case.pdf"
    c = canvas.Canvas(str(pdf_file))
    c.drawString(50, 800, "Identificador del caso: R-055429")
    c.drawString(50, 780, "Módulo: Pasarela")
    c.drawString(50, 760, "Título: Usuario para Reportería")
    c.drawString(50, 740, "Objetivo: Validar creación de usuario para reportería")
    c.drawString(50, 720, "Link: https://qa.example.test/login")
    c.drawString(50, 700, "Ambiente: QA")
    c.drawString(50, 680, "Analista: QA Tester")
    c.drawString(50, 660, "Usuario: test_user")
    c.drawString(50, 640, "Clave: test_password")
    
    c.drawString(50, 600, 'Step 1: (9/2/2026 3:21:33 PM) User left click on "Mantenimiento (link)" in "Browser"')
    c.drawString(50, 580, 'Step 2: (9/2/2026 3:21:34 PM) User left click on "Seguridad (link)" in "Browser"')
    c.drawString(50, 560, 'Step 3: (9/2/2026 3:21:38 PM) User left click on "Agregar Usuario (button)" in "Browser"')
    c.save()
    return pdf_file


def test_docx_parser_with_sample_fixture(sample_docx_path: Path):
    """Verifies DocxParser extracts metadata and steps accurately from a fixture."""
    parser = DocxParser()
    test_case = parser.parse(sample_docx_path)
    
    assert isinstance(test_case, TestCase)
    assert test_case.case_id == "R-055429"
    assert test_case.module == "Pasarela"
    assert "Usuario para Reportería" in test_case.title
    assert "creación y edición de usuarios" in test_case.objective
    assert test_case.credentials.get("username") == "test_user"
    assert test_case.credentials.get("password") == "test_password"
    assert test_case.url == "https://qa.example.test/login"
    assert test_case.environment == "QA"
    assert test_case.developer == "QA Tester"
    
    assert len(test_case.steps) == 3
    assert test_case.steps[0].step_num == 2
    assert test_case.steps[0].action == "fill"
    assert test_case.steps[1].step_num == 3
    assert test_case.steps[1].action == "click"
    assert test_case.steps[1].target == "Mantenimiento (link)"
    assert test_case.steps[2].step_num == 4
    assert test_case.steps[2].action == "click"
    assert test_case.steps[2].target == "Seguridad (link)"


def test_docx_parser_with_production_document():
    """Verifies DocxParser extracts metadata and all steps from the real production docx."""
    if not PROD_DOCX_PATH.exists():
        pytest.skip(f"Production document {PROD_DOCX_PATH} not found.")
        
    parser = DocxParser()
    test_case = parser.parse(PROD_DOCX_PATH)
    
    assert isinstance(test_case, TestCase)
    assert test_case.case_id == "R-055429"
    assert test_case.module == "Pasarela"
    assert "Usuario para Reportería" in test_case.title
    assert "Verificar el funcionamiento" in test_case.objective
    assert bool(test_case.credentials.get("username"))
    assert bool(test_case.credentials.get("password"))
    assert bool(test_case.url and test_case.url.startswith("http"))
    assert test_case.environment == "QA"
    assert bool(test_case.developer)
    
    # The production document contains 170 action steps
    assert len(test_case.steps) == 170
    assert all(isinstance(s, StepDefinition) for s in test_case.steps)
    assert test_case.steps[0].step_num == 2
    assert test_case.steps[0].action == "fill"
    assert test_case.steps[1].step_num == 3
    assert test_case.steps[1].action == "click"
    assert test_case.steps[1].target == "Mantenimiento (link)"


def test_pdf_parser_with_sample_fixture(sample_pdf_path: Path):
    """Verifies PdfParser extracts metadata and steps accurately from a sample PDF."""
    parser = PdfParser()
    test_case = parser.parse(sample_pdf_path)
    
    assert isinstance(test_case, TestCase)
    assert test_case.case_id == "R-055429"
    assert test_case.module == "Pasarela"
    assert test_case.title == "Usuario para Reportería"
    assert "Validar creación de usuario" in test_case.objective
    assert test_case.credentials.get("username") == "test_user"
    assert test_case.credentials.get("password") == "test_password"
    assert test_case.url == "https://qa.example.test/login"
    assert test_case.environment == "QA"
    assert test_case.developer == "QA Tester"
    
    assert len(test_case.steps) == 3
    assert test_case.steps[0].step_num == 1
    assert test_case.steps[0].action == "click"
    assert test_case.steps[0].target == "Mantenimiento (link)"
    assert test_case.steps[1].step_num == 2
    assert test_case.steps[1].action == "click"
    assert test_case.steps[1].target == "Seguridad (link)"
    assert test_case.steps[2].step_num == 3
    assert test_case.steps[2].action == "click"
    assert test_case.steps[2].target == "Agregar Usuario (button)"


def test_unified_doc_parser_dispatch(sample_docx_path: Path, sample_pdf_path: Path, tmp_path: Path):
    """Verifies DocParser.parse dispatches correctly to DocxParser and PdfParser based on file extension."""
    docx_case = DocParser.parse(sample_docx_path)
    assert isinstance(docx_case, TestCase)
    assert docx_case.case_id == "R-055429"
    
    pdf_case = DocParser.parse(sample_pdf_path)
    assert isinstance(pdf_case, TestCase)
    assert pdf_case.case_id == "R-055429"
    
    unsupported_file = tmp_path / "test.unsupported"
    unsupported_file.write_text("dummy")
    with pytest.raises(ValueError, match="Unsupported file format"):
        DocParser.parse(unsupported_file)
