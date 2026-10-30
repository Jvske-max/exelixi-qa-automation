"""Unit tests for Normalizer Providers and E2E Pipeline (T8)."""

from __future__ import annotations

import json
import socket
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock, patch

import docx
import pytest
import requests

from normalizer.cli import main as cli_main
from normalizer.grounding import check as check_grounding
from normalizer.providers import (
    AgentFileProvider,
    ApiProvider,
    NormalizerProvider,
    get_provider,
)
from normalizer.providers.masking import mask_dump_for_provider
from normalizer.schema import Case, NormalizedDoc, Step


@pytest.fixture(autouse=True)
def block_network(monkeypatch: pytest.MonkeyPatch) -> None:
    """Ensure zero real network connections are made across unit tests."""

    def forbidden_socket(*args: Any, **kwargs: Any) -> Any:
        raise RuntimeError("Real network access forbidden in unit tests")

    monkeypatch.setattr(socket, "create_connection", forbidden_socket)


@pytest.fixture
def sample_raw_dump() -> dict[str, Any]:
    """Provide a raw extracted dump containing diverse sensitive data."""
    return {
        "meta": {
            "módulo": "Seguridad",
            "identificador del caso": "R-PROV-01",
            "ambiente": "QA",
        },
        "paragraphs": [
            {
                "line": 1,
                "text": "Iniciamos sesión con usuario admin y clave: 'SuperSecret123!'",
                "images": [],
            },
            {
                "line": 2,
                "text": "Consultamos la cédula con puntos V-12.345.678 del titular.",
                "images": [],
            },
            {
                "line": 3,
                "text": "Luego validamos la cédula sin puntos V-12345678 en el sistema.",
                "images": [],
            },
            {
                "line": 4,
                "text": "Se envía notificación al correo analista.qa@example.test.",
                "images": [],
            },
            {
                "line": 5,
                "text": "Verificamos RIF empresarial J-12345678-0 de la compañía.",
                "images": [],
            },
            {
                "line": 6,
                "text": "El token de autorización es eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJzdWIiOiIxMjM0NTY3ODkwIn0.dozjgNryP4J3jVmNHl0w5N_XgL0n3I9PlF5",
                "images": [],
            },
            {
                "line": 7,
                "text": "Hacemos clic en el botón Continuar.",
                "images": [],
            },
        ],
        "cases": [
            {
                "n": 1,
                "title": "Validación de Datos Sensibles",
                "start_line": 1,
                "end_line": 7,
            }
        ],
        "links": ["https://qa.example.test/portal"],
    }


@pytest.fixture
def synthetic_sample_docx(tmp_path: Path) -> Path:
    """Create a synthetic docx specification for provider E2E tests."""
    doc_path = tmp_path / "synthetic_spec.docx"
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=4)
    table.rows[0].cells[0].text = "Fecha"
    table.rows[0].cells[1].text = "Módulo"
    table.rows[0].cells[2].text = "Identificador del caso"
    table.rows[0].cells[3].text = "Ambiente"
    table.rows[1].cells[0].text = "2026-10-30"
    table.rows[1].cells[1].text = "Seguridad"
    table.rows[1].cells[2].text = "R-SYNTH-PROV"
    table.rows[1].cells[3].text = "QA"

    doc.add_paragraph("Link: https://qa.example.test/portal")
    doc.add_paragraph("Caso 1: Flujo de Autenticación")
    doc.add_paragraph("Objetivo de la Prueba: Validar ingreso de usuario")
    doc.add_paragraph("Iniciamos sesión con usuario ${QA_USER_ADMIN} y clave ${QA_PASS_ADMIN}")
    doc.add_paragraph("Paso 1: clic en Iniciar Sesión")
    doc.save(str(doc_path))
    return doc_path


@pytest.mark.smoke
@pytest.mark.api
def test_mask_dump_for_provider_all_sensitive_types(sample_raw_dump: dict[str, Any]) -> None:
    """Positive and negative test for deterministic masking of all sensitive data types."""
    # Positive case: all sensitive data types are masked with deterministic markers
    masked_dump, mapping = mask_dump_for_provider(sample_raw_dump)

    all_masked_text = " ".join(p["text"] for p in masked_dump["paragraphs"])

    # Verify sequential deterministic markers are present
    assert "[PASS_1]" in all_masked_text
    assert "[CEDULA_1]" in all_masked_text
    assert "[EMAIL_1]" in all_masked_text
    assert "[RIF_1]" in all_masked_text
    assert "[JWT_1]" in all_masked_text

    # Verify original sensitive data does NOT leak anywhere in the masked dump
    assert "SuperSecret123!" not in all_masked_text
    assert "12.345.678" not in all_masked_text
    assert "12345678" not in all_masked_text
    assert "analista.qa@example.test" not in all_masked_text
    assert "J-12345678-0" not in all_masked_text
    assert "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9" not in all_masked_text

    # Verify mapping records markers
    assert "[PASS_1]" in mapping
    assert "[EMAIL_1]" in mapping
    assert mapping["[PASS_1]"] == "SuperSecret123!"
    assert mapping["[EMAIL_1]"] == "analista.qa@example.test"

    # Negative case: non-sensitive text and environment variables must NOT be masked
    assert "Iniciamos sesión con usuario admin" in all_masked_text
    assert "Hacemos clic en el botón Continuar." in all_masked_text
    raw_env_dump = {
        "meta": {"ambiente": "QA"},
        "paragraphs": [
            {"line": 1, "text": "Usamos clave ${QA_PASS_ADMIN} y rol ${QA_USER_ADMIN}", "images": []}
        ],
        "cases": [],
        "links": [],
    }
    env_masked, env_mapping = mask_dump_for_provider(raw_env_dump)
    assert "${QA_PASS_ADMIN}" in env_masked["paragraphs"][0]["text"]
    assert len(env_mapping) == 0


@pytest.mark.smoke
@pytest.mark.api
def test_masking_invariance_for_api_and_grounding(sample_raw_dump: dict[str, Any]) -> None:
    """Verify invariant: text sent to API and text checked by grounding share the exact same masking."""
    # Positive case: exact same masked dump is produced deterministically across multiple calls
    masked_1, map_1 = mask_dump_for_provider(sample_raw_dump)
    masked_2, map_2 = mask_dump_for_provider(sample_raw_dump)

    assert masked_1 == masked_2
    assert map_1 == map_2

    # A normalized doc quoting the masked text must pass grounding check against masked_dump
    valid_norm = NormalizedDoc(
        meta={"identificador del caso": "R-PROV-01", "módulo": "Seguridad"},
        cases=[
            Case(
                n=1,
                title="Validación de Datos Sensibles",
                objective="Verificar datos",
                url="https://qa.example.test/portal",
                credentials={"usuario": "${QA_USER_ADMIN}", "clave": "${QA_PASS_ADMIN}"},
                steps=[
                    Step(
                        action="click",
                        target="Continuar",
                        source_quote="Hacemos clic en el botón Continuar.",
                        source_ref="Línea 7",
                    )
                ],
            )
        ],
    )
    violations = check_grounding(valid_norm, masked_1)
    assert violations == []

    # Negative case: doc quoting unmasked real sensitive text produces grounding violations
    invalid_norm = NormalizedDoc(
        meta={"identificador del caso": "R-PROV-01", "módulo": "Seguridad"},
        cases=[
            Case(
                n=1,
                title="Validación de Datos Sensibles",
                objective="Verificar datos",
                url="https://qa.example.test/portal",
                credentials={"usuario": "${QA_USER_ADMIN}", "clave": "${QA_PASS_ADMIN}"},
                steps=[
                    Step(
                        action="fill",
                        target="SuperSecret123!",
                        source_quote="Iniciamos sesión con usuario admin y clave: 'SuperSecret123!'",
                        source_ref="Línea 1",
                    )
                ],
            )
        ],
    )
    violations_invalid = check_grounding(invalid_norm, masked_1)
    assert len(violations_invalid) > 0
    assert any(v.type == "ungrounded_quote" for v in violations_invalid)


@pytest.mark.smoke
@pytest.mark.api
def test_api_provider_env_validation_and_no_leak(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify ApiProvider validates env vars, enforces https, and never leaks keys or URLs in errors."""
    monkeypatch.delenv("NORMALIZER_API_KEY", raising=False)
    monkeypatch.delenv("NORMALIZER_API_URL", raising=False)

    # Negative case 1: Missing API KEY raises ValueError without leaking
    with pytest.raises(ValueError) as exc_info:
        ApiProvider()
    err_msg = str(exc_info.value)
    assert "NORMALIZER_API_KEY" in err_msg

    # Negative case 2: Missing API URL raises ValueError
    monkeypatch.setenv("NORMALIZER_API_KEY", "secret-test-key-999")
    with pytest.raises(ValueError) as exc_info:
        ApiProvider()
    assert "NORMALIZER_API_URL" in str(exc_info.value)

    # Negative case 3: Non-https URL scheme raises ValueError and does NOT leak URL or secret
    insecure_url = "http://insecure-api.internal.test/v1/normalize"
    secret_key = "super-secret-key-alpha-bravo"
    monkeypatch.setenv("NORMALIZER_API_URL", insecure_url)
    monkeypatch.setenv("NORMALIZER_API_KEY", secret_key)

    with pytest.raises(ValueError) as exc_info:
        ApiProvider()
    err_insecure = str(exc_info.value)
    assert "https" in err_insecure
    # CRITICAL: Secret key and full url must NEVER be reflected in error message
    assert secret_key not in err_insecure
    assert insecure_url not in err_insecure

    # Positive case: Valid https URL and key initialize cleanly
    valid_url = "https://ai.example.test/v1/normalize"
    monkeypatch.setenv("NORMALIZER_API_URL", valid_url)
    provider = ApiProvider()
    assert isinstance(provider, NormalizerProvider)
    assert provider.max_chars == 20000
    assert provider.max_images == 5
    assert provider.max_retries == 3


@pytest.mark.smoke
@pytest.mark.api
def test_api_provider_dry_run(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify --dry-run writes masked payload to build/<id>/dry_run_payload.json without network calls."""
    monkeypatch.setenv("NORMALIZER_API_URL", "https://ai.example.test/v1/normalize")
    monkeypatch.setenv("NORMALIZER_API_KEY", "fake-key")

    provider = ApiProvider()
    custom_dry_run_path = tmp_path / "build" / "R-DRY-01" / "dry_run_payload.json"

    raw_dump = {
        "meta": {"identificador del caso": "R-DRY-01", "módulo": "Ventas"},
        "paragraphs": [
            {"line": 1, "text": "Clave: 'MySecretPass'", "images": []},
            {"line": 2, "text": "Paso 1: clic en Enviar", "images": []},
        ],
        "cases": [{"n": 1, "title": "Dry Run Test", "start_line": 1, "end_line": 2}],
        "links": ["https://qa.example.test/portal"],
    }

    # Positive case: normalize with dry_run=True creates payload file on disk
    res = provider.normalize(
        dump=raw_dump,
        dry_run=True,
        dry_run_path=custom_dry_run_path,
    )
    assert custom_dry_run_path.exists()
    payload_content = json.loads(custom_dry_run_path.read_text(encoding="utf-8"))

    # Masked payload contains [PASS_1] and zero real secret text
    assert "dump" in payload_content
    paragraphs_str = json.dumps(payload_content["dump"]["paragraphs"])
    assert "[PASS_1]" in paragraphs_str
    assert "MySecretPass" not in paragraphs_str

    # Negative case: verify normalize returns without making any real network socket connection
    assert res is not None


@pytest.mark.smoke
@pytest.mark.api
def test_api_provider_with_images_warning(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """Verify warning printed to sys.stderr only when with_images=True."""
    monkeypatch.setenv("NORMALIZER_API_URL", "https://ai.example.test/v1/normalize")
    monkeypatch.setenv("NORMALIZER_API_KEY", "fake-key")

    provider = ApiProvider()
    raw_dump = {
        "meta": {"identificador del caso": "R-IMG-01"},
        "paragraphs": [{"line": 1, "text": "Paso 1: clic", "images": []}],
        "cases": [],
        "links": [],
    }

    # Positive case: with_images=True prints explicit warning to sys.stderr
    provider.normalize(dump=raw_dump, with_images=True, dry_run=True)
    captured = capsys.readouterr()
    expected_warning = (
        "ADVERTENCIA: Las capturas de pantalla no admiten enmascaramiento automático "
        "y pueden contener datos visibles."
    )
    assert expected_warning in captured.err

    # Negative case: with_images=False does NOT print screenshot warning
    provider.normalize(dump=raw_dump, with_images=False, dry_run=True)
    captured_no_img = capsys.readouterr()
    assert expected_warning not in captured_no_img.err


@pytest.mark.smoke
@pytest.mark.api
def test_api_provider_fail_safe_on_timeout(
    synthetic_sample_docx: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Verify fail-safe: network timeout causes exit code != 0, 0 docx rendered, and never READY."""
    monkeypatch.setenv("NORMALIZER_API_URL", "https://ai.example.test/v1/normalize")
    monkeypatch.setenv("NORMALIZER_API_KEY", "fake-key")

    # Simulate network timeout during requests.post
    with patch("requests.Session.post", side_effect=requests.exceptions.Timeout("Connection timed out")):
        out_dir = tmp_path / "timeout_output"
        code = cli_main(
            [
                "normalize",
                str(synthetic_sample_docx),
                "--provider",
                "api",
                "--out",
                str(out_dir),
            ]
        )

        # Positive verification: execution terminates with error code != 0
        assert code != 0

        # Negative verification: ZERO normalized docx files exist on disk, never READY
        rendered_files = list(out_dir.glob("*.docx")) if out_dir.exists() else []
        assert len(rendered_files) == 0


@pytest.mark.smoke
@pytest.mark.api
def test_api_provider_fail_safe_on_invalid_response(
    synthetic_sample_docx: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Verify fail-safe: provider returning invalid schema or ungrounded response is rejected."""
    monkeypatch.setenv("NORMALIZER_API_URL", "https://ai.example.test/v1/normalize")
    monkeypatch.setenv("NORMALIZER_API_KEY", "fake-key")

    # Mock response returning hallucinated/ungrounded quote
    mock_resp = MagicMock()
    mock_resp.status_code = 200
    mock_resp.json.return_value = {
        "meta": {"módulo": "Seguridad", "identificador del caso": "R-SYNTH-PROV"},
        "cases": [
            {
                "n": 1,
                "title": "Caso Alucinado",
                "objective": "No grounded",
                "url": "https://qa.example.test/portal",
                "credentials": {"usuario": "${QA_USER_ADMIN}", "clave": "${QA_PASS_ADMIN}"},
                "steps": [
                    {
                        "action": "click",
                        "target": "Boton Inexistente",
                        "source_quote": "Esta cita inventada jamas existio en el documento original",
                        "source_ref": "Línea 999",
                    }
                ],
            }
        ],
    }

    with patch("requests.Session.post", return_value=mock_resp):
        out_dir = tmp_path / "invalid_output"
        code = cli_main(
            [
                "normalize",
                str(synthetic_sample_docx),
                "--provider",
                "api",
                "--out",
                str(out_dir),
            ]
        )

        # Positive verification: rejected with non-zero exit code
        assert code != 0

        # Negative verification: ZERO docx rendered, never marked READY
        docx_files = list(out_dir.glob("*.docx")) if out_dir.exists() else []
        assert len(docx_files) == 0


@pytest.mark.smoke
@pytest.mark.api
def test_agent_file_provider_default(
    synthetic_sample_docx: Path, tmp_path: Path
) -> None:
    """Positive and negative test for default agent_file provider reading normalized/<stem>.json."""
    # Setup pre-generated normalized file in temporary normalized/ directory
    norm_dir = tmp_path / "normalized"
    norm_dir.mkdir(parents=True, exist_ok=True)
    stem = synthetic_sample_docx.stem

    valid_norm_doc = NormalizedDoc(
        meta={"módulo": "Seguridad", "identificador del caso": "R-SYNTH-PROV"},
        cases=[
            Case(
                n=1,
                title="Flujo de Autenticación",
                objective="Validar ingreso de usuario",
                url="https://qa.example.test/portal",
                credentials={"usuario": "${QA_USER_ADMIN}", "clave": "${QA_PASS_ADMIN}"},
                steps=[
                    Step(
                        action="click",
                        target="Iniciar Sesión",
                        source_quote="Paso 1: clic en Iniciar Sesión",
                        source_ref="Línea 5",
                    )
                ],
            )
        ],
    )
    (norm_dir / f"{stem}.json").write_text(valid_norm_doc.model_dump_json(indent=2), encoding="utf-8")

    out_dir = tmp_path / "rendered_canonical"

    # Negative CLI case: file not present in default normalized/ returns code != 0
    code_missing = cli_main(
        [
            "normalize",
            str(synthetic_sample_docx),
            "--out",
            str(out_dir),
        ]
    )
    assert code_missing != 0

    # Positive CLI case: when normalized/<stem>.json is present, normalizes and renders cleanly
    default_norm_file = Path("normalized") / f"{stem}.json"
    try:
        default_norm_file.write_text(valid_norm_doc.model_dump_json(indent=2), encoding="utf-8")
        code = cli_main(
            [
                "normalize",
                str(synthetic_sample_docx),
                "--out",
                str(out_dir),
            ]
        )
        assert code == 0
        rendered_docx = out_dir / "R-SYNTH-PROV-caso1.docx"
        assert rendered_docx.exists()
    finally:
        if default_norm_file.exists():
            default_norm_file.unlink()

    # Direct provider tests with custom directory
    provider = AgentFileProvider(normalized_dir=norm_dir, stem=stem)
    raw_dump = {
        "meta": {"identificador del caso": "R-SYNTH-PROV"},
        "paragraphs": [{"line": 1, "text": "Paso 1: clic en Iniciar Sesión", "images": []}],
        "cases": [],
        "links": [],
    }
    loaded = provider.normalize(raw_dump)
    assert loaded["meta"]["identificador del caso"] == "R-SYNTH-PROV"
    assert len(loaded["cases"]) == 1

    # Negative provider case: non-existent stem raises FileNotFoundError
    missing_provider = AgentFileProvider(normalized_dir=norm_dir, stem="non_existent_stem")
    with pytest.raises(FileNotFoundError):
        missing_provider.normalize(raw_dump)


@pytest.mark.smoke
@pytest.mark.api
def test_get_provider_factory() -> None:
    """Positive and negative test for get_provider factory."""
    # Positive case 1: get_provider('agent_file') returns AgentFileProvider
    p_file = get_provider("agent_file")
    assert isinstance(p_file, AgentFileProvider)

    # Positive case 2: get_provider('api') returns ApiProvider when env configured
    with patch.dict(
        "os.environ",
        {
            "NORMALIZER_API_URL": "https://ai.example.test/v1/normalize",
            "NORMALIZER_API_KEY": "fake-key",
        },
    ):
        p_api = get_provider("api")
        assert isinstance(p_api, ApiProvider)

    # Negative case: unknown provider raises ValueError
    with pytest.raises(ValueError) as exc:
        get_provider("unsupported_provider_xyz")
    assert "Unknown provider" in str(exc.value)
