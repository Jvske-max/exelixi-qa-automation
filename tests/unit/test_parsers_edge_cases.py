"""Edge-case unit tests for document parsers (DocParser, DocxParser, PdfParser)."""

from pathlib import Path

import docx
import pytest
from reportlab.pdfgen import canvas

from parsers import DocParser, DocxParser, PdfParser
from parsers.models import TestCase

TEST_DUMMY_URL = "https://example.com/test-login"


# ==============================================================================
# Requirement 1: Non-existent file path tests (FileNotFoundError)
# ==============================================================================


@pytest.mark.smoke
def test_non_existent_pdf_file_raises_file_not_found(tmp_path: Path) -> None:
    """Negative: verifies DocParser and PdfParser raise FileNotFoundError for missing PDF files."""
    missing_pdf = tmp_path / "non_existent_document.pdf"

    with pytest.raises(FileNotFoundError):
        DocParser.parse(missing_pdf)

    with pytest.raises(FileNotFoundError):
        PdfParser.parse(missing_pdf)

    with pytest.raises(FileNotFoundError):
        DocParser.parse(str(missing_pdf))


@pytest.mark.smoke
def test_existing_file_does_not_raise_file_not_found(tmp_path: Path) -> None:
    """Positive: verifies parsing existing files succeeds without FileNotFoundError."""
    existing_pdf = tmp_path / "valid_document.pdf"
    c = canvas.Canvas(str(existing_pdf))
    c.drawString(50, 750, "Caso ID: R-101010")
    c.drawString(50, 730, 'Step 1: User left click on "Save" in "App"')
    c.save()

    test_case = DocParser.parse(existing_pdf)
    assert isinstance(test_case, TestCase)
    assert test_case.case_id == "R-101010"
    assert len(test_case.steps) == 1


# ==============================================================================
# Requirement 2: Unsupported file extensions tests (ValueError)
# ==============================================================================


@pytest.mark.smoke
@pytest.mark.parametrize(
    "extension",
    [".xlsx", ".txt", ".csv", ".json", ".xml", ".png", ""],
)
def test_unsupported_file_extension_raises_value_error(tmp_path: Path, extension: str) -> None:
    """Negative: verifies DocParser.parse raises ValueError for unsupported extensions."""
    file_name = f"document{extension}" if extension else "document_without_ext"
    unsupported_file = tmp_path / file_name
    unsupported_file.write_text("dummy test content", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file format"):
        DocParser.parse(unsupported_file)


@pytest.mark.smoke
def test_supported_file_extensions_do_not_raise_value_error(tmp_path: Path) -> None:
    """Positive: verifies supported extensions (.docx, .pdf) are accepted by DocParser."""
    valid_docx = tmp_path / "supported.docx"
    doc = docx.Document()
    doc.add_paragraph("Step 1: User left click on \"Button\" in \"App\"")
    doc.save(str(valid_docx))

    test_case_docx = DocParser.parse(valid_docx)
    assert isinstance(test_case_docx, TestCase)

    valid_pdf = tmp_path / "supported.pdf"
    c = canvas.Canvas(str(valid_pdf))
    c.drawString(50, 750, 'Step 1: User left click on "Button" in "App"')
    c.save()

    test_case_pdf = DocParser.parse(valid_pdf)
    assert isinstance(test_case_pdf, TestCase)


# ==============================================================================
# Requirement 3: Missing tables and missing metadata fallback defaults
# ==============================================================================


@pytest.mark.smoke
def test_docx_missing_tables_and_metadata_falls_back_to_defaults(tmp_path: Path) -> None:
    """Negative/Edge: verifies DocxParser falls back gracefully when tables and metadata are absent."""
    doc_path = tmp_path / "plain_case.docx"
    doc = docx.Document()
    doc.add_paragraph('Step 1: (9/2/2026) User left click on "Aceptar (button)" in "App"')
    doc.save(str(doc_path))

    test_case = DocxParser.parse(doc_path)

    assert isinstance(test_case, TestCase)
    assert test_case.case_id == "UNKNOWN"
    assert test_case.module == "General"
    assert test_case.title == "plain_case"
    assert test_case.objective == ""
    assert test_case.environment == "QA"
    assert test_case.developer == "Exelixi QA"
    assert test_case.credentials == {}
    assert test_case.url is None
    assert len(test_case.steps) == 1


@pytest.mark.smoke
def test_docx_malformed_and_empty_tables_fallback(tmp_path: Path) -> None:
    """Negative/Edge: verifies DocxParser handles empty and malformed tables without crashing."""
    doc_path = tmp_path / "malformed_tables.docx"
    doc = docx.Document()

    # Empty table with 0 rows
    doc.add_table(rows=0, cols=0)

    # 1-row table with empty cells
    t1 = doc.add_table(rows=1, cols=2)
    t1.rows[0].cells[0].text = ""
    t1.rows[0].cells[1].text = ""

    # 2-row table where header and value are identical or empty
    t2 = doc.add_table(rows=2, cols=2)
    t2.rows[0].cells[0].text = "Header"
    t2.rows[0].cells[1].text = ""
    t2.rows[1].cells[0].text = "Header"
    t2.rows[1].cells[1].text = ""

    doc.add_paragraph('Step 1: User left click on "Close"')
    doc.save(str(doc_path))

    test_case = DocxParser.parse(doc_path)
    assert test_case.case_id == "UNKNOWN"
    assert test_case.module == "General"
    assert test_case.developer == "Exelixi QA"
    assert len(test_case.steps) == 1


@pytest.mark.smoke
def test_pdf_missing_metadata_falls_back_to_defaults(tmp_path: Path) -> None:
    """Negative/Edge: verifies PdfParser falls back to default values when metadata fields are missing."""
    pdf_path = tmp_path / "plain_pdf.pdf"
    c = canvas.Canvas(str(pdf_path))
    c.drawString(50, 700, 'Step 1: User left click on "Submit" in "App"')
    c.save()

    test_case = PdfParser.parse(pdf_path)

    assert isinstance(test_case, TestCase)
    assert test_case.case_id == "UNKNOWN"
    assert test_case.module == "General"
    assert test_case.title == "plain_pdf"
    assert test_case.objective == ""
    assert test_case.environment == "QA"
    assert test_case.developer == "Exelixi QA"
    assert test_case.credentials == {}
    assert test_case.url is None
    assert len(test_case.steps) == 1


@pytest.mark.smoke
def test_case_id_extracted_from_filename_when_missing_in_content(tmp_path: Path) -> None:
    """Edge: verifies case_id is parsed from filename pattern (R-XXXXXX) when missing in body."""
    docx_path = tmp_path / "spec_R-884422_draft.docx"
    doc = docx.Document()
    doc.add_paragraph('Step 1: User left click on "Menu"')
    doc.save(str(docx_path))

    test_case_docx = DocxParser.parse(docx_path)
    assert test_case_docx.case_id == "R-884422"

    pdf_path = tmp_path / "spec_R-773311_draft.pdf"
    c = canvas.Canvas(str(pdf_path))
    c.drawString(50, 700, 'Step 1: User left click on "Menu"')
    c.save()

    test_case_pdf = PdfParser.parse(pdf_path)
    assert test_case_pdf.case_id == "R-773311"


@pytest.mark.smoke
def test_document_with_full_metadata_populates_explicit_fields(tmp_path: Path) -> None:
    """Positive: verifies fully specified documents populate all fields without using fallback defaults."""
    doc_path = tmp_path / "full_meta.docx"
    doc = docx.Document()

    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Módulo"
    table.rows[0].cells[1].text = "Billing"
    table.rows[1].cells[0].text = "Ambiente de prueba"
    table.rows[1].cells[1].text = "Staging"
    table.rows[2].cells[0].text = "Analista Programador"
    table.rows[2].cells[1].text = "QA Lead"

    doc.add_paragraph("Caso # 42: Facturación Electrónica")
    doc.add_paragraph("Objetivo de la Prueba: Validar ciclo de cobro")
    doc.add_paragraph(f"Link: {TEST_DUMMY_URL}")
    doc.add_paragraph("Usuario: test_analyst")
    doc.add_paragraph("Clave: test_secret_123")
    doc.add_paragraph('Step 1: User left click on "Pagar"')
    doc.save(str(doc_path))

    test_case = DocxParser.parse(doc_path)
    assert test_case.module == "Billing"
    assert test_case.environment == "Staging"
    assert test_case.developer == "QA Lead"
    assert "Facturación Electrónica" in test_case.title
    assert test_case.objective == "Validar ciclo de cobro"
    assert test_case.url == TEST_DUMMY_URL
    assert test_case.credentials.get("username") == "test_analyst"
    assert test_case.credentials.get("password") == "test_secret_123"


# ==============================================================================
# Requirement 4: Unusual and multiline keyboard input parsing
# ==============================================================================


@pytest.mark.smoke
def test_psr_step_parsing_with_unusual_keyboard_inputs(tmp_path: Path) -> None:
    """Edge: verifies PSR step parsing with shortcut sequences, symbols, and escaped multiline inputs."""
    doc_path = tmp_path / "unusual_inputs.docx"
    doc = docx.Document()

    # Step with hotkey sequences (CTRL-V, TABULACION, ENTRAR)
    doc.add_paragraph(
        'Step 1: (9/2/2026 3:21:29 PM) User keyboard input on '
        '"Mundial (window)" in "App" [... CTRL-V TABULACION ... CTRL-V ENTRAR ENTRAR]'
    )
    # Step with special symbols, emails, and punctuation
    doc.add_paragraph(
        'Step 2: (9/2/2026 3:21:30 PM) User keyboard input on '
        '"Email (edit)" in "App" [qa_engineer+sub1@domain-test.org: !#$%&*()_+=]'
    )
    # Step with escaped multiline representation
    doc.add_paragraph(
        'Step 3: (9/2/2026 3:21:31 PM) User keyboard input on '
        '"Notes (edit)" in "App" [Line 1\\nLine 2\\nLine 3]'
    )
    # Step with keyboard input action without bracketed value
    doc.add_paragraph(
        'Step 4: (9/2/2026 3:21:32 PM) User keyboard input on "Next (button)" in "App"'
    )
    doc.save(str(doc_path))

    test_case = DocxParser.parse(doc_path)
    assert len(test_case.steps) == 4

    step1 = test_case.steps[0]
    assert step1.step_num == 1
    assert step1.action == "fill"
    assert step1.target == "Mundial (window)"
    assert step1.value == "... CTRL-V TABULACION ... CTRL-V ENTRAR ENTRAR"

    step2 = test_case.steps[1]
    assert step2.step_num == 2
    assert step2.action == "fill"
    assert step2.value == "qa_engineer+sub1@domain-test.org: !#$%&*()_+="

    step3 = test_case.steps[2]
    assert step3.step_num == 3
    assert step3.action == "fill"
    assert step3.value == "Line 1\\nLine 2\\nLine 3"

    step4 = test_case.steps[3]
    assert step4.step_num == 4
    assert step4.action == "fill"
    assert step4.value is None


@pytest.mark.smoke
def test_standard_step_parsing_with_unusual_and_multiline_inputs(tmp_path: Path) -> None:
    """Edge: verifies standard step parsing (Step/Paso) with multiline strings and symbols."""
    doc_path = tmp_path / "standard_inputs.docx"
    doc = docx.Document()

    doc.add_paragraph('Step 1: Fill "First line\\nSecond line\\nThird line" in "Address Textarea"')
    doc.add_paragraph('Paso 2: Escribir "secure#pass!_123" en "Clave de Acceso"')
    doc.add_paragraph('Paso 3: Navegar a "https://example.com/portal"')
    doc.save(str(doc_path))

    test_case = DocxParser.parse(doc_path)
    assert len(test_case.steps) == 3

    assert test_case.steps[0].action == "fill"
    assert test_case.steps[0].target == "Address Textarea"
    assert test_case.steps[0].value == "First line\\nSecond line\\nThird line"

    assert test_case.steps[1].action == "fill"
    assert test_case.steps[1].target == "Clave de Acceso"
    assert test_case.steps[1].value == "secure#pass!_123"

    assert test_case.steps[2].action == "navigate"
    assert test_case.steps[2].target == "https://example.com/portal"


@pytest.mark.smoke
def test_positive_standard_keyboard_input_step(tmp_path: Path) -> None:
    """Positive: verifies regular single-line keyboard input is parsed cleanly."""
    doc_path = tmp_path / "normal_input.docx"
    doc = docx.Document()
    doc.add_paragraph('Step 1: (9/2/2026 3:21:29 PM) User keyboard input on "User (edit)" in "App" [johndoe]')
    doc.save(str(doc_path))

    test_case = DocxParser.parse(doc_path)
    assert len(test_case.steps) == 1
    assert test_case.steps[0].action == "fill"
    assert test_case.steps[0].target == "User (edit)"
    assert test_case.steps[0].value == "johndoe"


# ==============================================================================
# Requirement 5: Empty document handling (0 steps)
# ==============================================================================


@pytest.mark.smoke
def test_completely_empty_docx_document(tmp_path: Path) -> None:
    """Negative/Edge: verifies completely blank .docx returns TestCase with 0 steps without errors."""
    empty_docx = tmp_path / "empty_spec.docx"
    doc = docx.Document()
    doc.save(str(empty_docx))

    test_case = DocParser.parse(empty_docx)
    assert isinstance(test_case, TestCase)
    assert test_case.steps == []
    assert len(test_case.steps) == 0
    assert test_case.title == "empty_spec"
    assert test_case.case_id == "UNKNOWN"


@pytest.mark.smoke
def test_completely_empty_pdf_document(tmp_path: Path) -> None:
    """Negative/Edge: verifies completely blank .pdf returns TestCase with 0 steps without errors."""
    empty_pdf = tmp_path / "empty_spec.pdf"
    c = canvas.Canvas(str(empty_pdf))
    c.showPage()
    c.save()

    test_case = DocParser.parse(empty_pdf)
    assert isinstance(test_case, TestCase)
    assert test_case.steps == []
    assert len(test_case.steps) == 0
    assert test_case.title == "empty_spec"
    assert test_case.case_id == "UNKNOWN"


@pytest.mark.smoke
def test_document_with_metadata_only_has_zero_steps(tmp_path: Path) -> None:
    """Edge: verifies document having metadata but no actionable steps parses 0 steps."""
    doc_path = tmp_path / "metadata_only.docx"
    doc = docx.Document()
    doc.add_paragraph("Caso # 10: Escenario sin pasos")
    doc.add_paragraph("Objetivo: Comprobar documento sin acciones")
    doc.add_paragraph(f"Link: {TEST_DUMMY_URL}")
    doc.save(str(doc_path))

    test_case = DocParser.parse(doc_path)
    assert isinstance(test_case, TestCase)
    assert len(test_case.steps) == 0
    assert test_case.steps == []
    assert "Escenario sin pasos" in test_case.title
    assert test_case.objective == "Comprobar documento sin acciones"


@pytest.mark.smoke
def test_positive_document_with_steps_has_non_zero_steps(tmp_path: Path) -> None:
    """Positive: verifies normal document contains > 0 parsed steps."""
    doc_path = tmp_path / "active_spec.docx"
    doc = docx.Document()
    doc.add_paragraph('Step 1: User left click on "Iniciar"')
    doc.add_paragraph('Step 2: User left click on "Finalizar"')
    doc.save(str(doc_path))

    test_case = DocParser.parse(doc_path)
    assert len(test_case.steps) == 2
    assert test_case.steps[0].action == "click"
    assert test_case.steps[1].action == "click"
