"""Integration tests for DocRunner execution engine, PDF reporting, and CLI."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import docx
import pypdf
import pytest

from parsers.models import StepDefinition, TestCase
from reporting.pdf_reporter import PDFReporter
from runners.doc_runner import DocRunner, main

MOCK_HTML_CONTENT = """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Integration Test Portal</title>
</head>
<body>
    <h1>QA Portal Test Page</h1>
    <form id="login-form">
        <div>
            <label for="username">Username</label>
            <input type="text" id="username" name="username" placeholder="Username" />
        </div>
        <div>
            <label for="password">Password</label>
            <input type="password" id="password" name="password" placeholder="Password" />
        </div>
        <button type="submit" id="login-button">Login</button>
    </form>
    <div id="welcome-message" style="display:none;">Welcome to Dashboard</div>
    <script>
        document.getElementById('login-form').addEventListener('submit', function(e) {
            e.preventDefault();
            document.getElementById('welcome-message').style.display = 'block';
        });
    </script>
</body>
</html>
"""


@pytest.fixture
def mock_html_page(tmp_path: Path) -> Path:
    """Creates a local mock HTML page for isolated UI tests."""
    html_file = tmp_path / "mock_portal.html"
    html_file.write_text(MOCK_HTML_CONTENT, encoding="utf-8")
    return html_file


@pytest.fixture
def mock_base_url(mock_html_page: Path) -> str:
    """Returns local file URL for the mock HTML page."""
    return f"file://{mock_html_page.resolve()}"


@pytest.fixture
def sample_docx_spec(tmp_path: Path, mock_base_url: str) -> Path:
    """Generates a valid .docx specification pointing to the mock portal."""
    doc_path = tmp_path / "login_spec.docx"
    doc = docx.Document()

    # Metadata table
    table = doc.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "Identificador del caso"
    table.rows[0].cells[1].text = "IT-001"
    table.rows[0].cells[2].text = "IT-001"

    table.rows[1].cells[0].text = "Módulo"
    table.rows[1].cells[1].text = "Security"
    table.rows[1].cells[2].text = "Security"

    table.rows[2].cells[0].text = "Ambiente de prueba"
    table.rows[2].cells[1].text = "IntegrationTest"
    table.rows[2].cells[2].text = "IntegrationTest"

    doc.add_paragraph(f"Link: {mock_base_url}")
    doc.add_paragraph("Caso # 1: Portal Login Flow")
    doc.add_paragraph("Objetivo de la Prueba")
    doc.add_paragraph("Verify full login interaction and reporting pipeline.")
    doc.add_paragraph("Usuario: test_user")
    doc.add_paragraph("Clave: test_pass")

    # Step paragraphs
    doc.add_paragraph('Step 1: (10/12/2026 10:00:00 AM) User left click on "portal (link)"')
    doc.add_paragraph(
        'Step 2: (10/12/2026 10:00:02 AM) User keyboard input on "Username" [... CTRL-V ...]'
    )
    doc.add_paragraph(
        'Step 3: (10/12/2026 10:00:04 AM) User keyboard input on "Password" [... CTRL-V ...]'
    )
    doc.add_paragraph(
        'Step 4: (10/12/2026 10:00:06 AM) User left click on "button#login-button (button)"'
    )

    doc.save(str(doc_path))
    return doc_path


@pytest.mark.e2e
def test_end_to_end_pipeline_success(
    tmp_path: Path, mock_base_url: str
) -> None:
    """Positive test: executes all steps against mock page and verifies PDF evidence."""
    screenshots_dir = tmp_path / "screenshots"
    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=5000,
        stop_on_failure=True,
    )

    test_case = TestCase(
        case_id="IT-001",
        module="Authentication",
        title="Valid Login Flow",
        objective="Verify complete form submission and screenshot generation.",
        environment="Integration",
        url=mock_base_url,
        credentials={"username": "qa_operator", "password": "SecurePassword123!"},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Navigate to mock portal",
                action="navigate",
                target=mock_base_url,
            ),
            StepDefinition(
                step_num=2,
                raw_text="Enter username",
                action="fill",
                target="input[name='username']",
                value="qa_operator",
            ),
            StepDefinition(
                step_num=3,
                raw_text="Enter password",
                action="fill",
                target="input[name='password']",
                value="SecurePassword123!",
            ),
            StepDefinition(
                step_num=4,
                raw_text="Click login button",
                action="click",
                target="button#login-button",
            ),
            StepDefinition(
                step_num=5,
                raw_text="Verify dashboard greeting appears",
                action="verify",
                target="Welcome to Dashboard",
            ),
        ],
    )

    report = runner.run_test_case(test_case=test_case, base_url=mock_base_url)

    # 1. Telemetry verification
    assert report.status == "PASSED"
    assert report.total_steps == 5
    assert report.passed_count == 5
    assert report.failed_count == 0
    assert report.duration_seconds > 0

    # 2. Screenshots verification
    captured_screenshots = list(screenshots_dir.glob("step_*.png"))
    assert len(captured_screenshots) == 5
    for step in report.steps:
        assert step.status == "PASSED"
        assert step.screenshot_path is not None
        assert Path(step.screenshot_path).exists()
        assert Path(step.screenshot_path).stat().st_size > 0

    # 3. PDF Evidence Report generation
    output_pdf = tmp_path / "reports" / "evidence_report_passed.pdf"
    reporter = PDFReporter()
    generated_path = reporter.generate_report(report, str(output_pdf))

    assert Path(generated_path).exists()
    assert Path(generated_path).stat().st_size > 1024

    pdf_reader = pypdf.PdfReader(generated_path)
    assert len(pdf_reader.pages) >= 1
    extracted_text = "".join(page.extract_text() or "" for page in pdf_reader.pages)
    assert "IT-001" in extracted_text
    assert "PASSED" in extracted_text
    assert "Welcome to Dashboard" in extracted_text


@pytest.mark.e2e
def test_pipeline_error_handling_failure(
    tmp_path: Path, mock_base_url: str
) -> None:
    """Negative test: catches element error, creates failure screenshot and FAILED PDF."""
    screenshots_dir = tmp_path / "screenshots_err"
    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=1000,
        stop_on_failure=True,
    )

    test_case = TestCase(
        case_id="IT-002",
        module="Authentication",
        title="Invalid Selector Failure Scenario",
        objective="Verify graceful failure handling and error capture.",
        environment="Integration",
        url=mock_base_url,
        credentials={},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Navigate to mock portal",
                action="navigate",
                target=mock_base_url,
            ),
            StepDefinition(
                step_num=2,
                raw_text="Click element that does not exist",
                action="click",
                target="#non-existent-button-selector-xyz",
            ),
            StepDefinition(
                step_num=3,
                raw_text="Subsequent step that should not run",
                action="click",
                target="button#login-button",
            ),
        ],
    )

    report = runner.run_test_case(test_case=test_case, base_url=mock_base_url)

    # 1. Telemetry and failure verification
    assert report.status == "FAILED"
    assert report.total_steps == 2  # Step 3 halted due to stop_on_failure
    assert report.passed_count == 1
    assert report.failed_count == 1
    assert report.steps[0].status == "PASSED"
    assert report.steps[1].status == "FAILED"
    assert "TimeoutError" in (report.steps[1].error_message or "")

    # 2. Failure screenshot verification
    failure_screenshot = screenshots_dir / "step_2_failed.png"
    assert failure_screenshot.exists()
    assert failure_screenshot.stat().st_size > 0
    assert report.steps[1].screenshot_path == str(failure_screenshot.resolve())

    # 3. PDF Evidence Report with FAILED badge
    output_pdf = tmp_path / "reports" / "evidence_report_failed.pdf"
    reporter = PDFReporter()
    generated_path = reporter.generate_report(report, str(output_pdf))

    assert Path(generated_path).exists()
    pdf_reader = pypdf.PdfReader(generated_path)
    extracted_text = "".join(page.extract_text() or "" for page in pdf_reader.pages)
    assert "IT-002" in extracted_text
    assert "FAILED" in extracted_text
    assert "TimeoutError" in extracted_text


@pytest.mark.smoke
def test_cli_help_flag() -> None:
    """Positive test: running CLI with --help flag succeeds with help text."""
    result = subprocess.run(
        [sys.executable, "-m", "runners.doc_runner", "--help"],
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0
    assert "Document-Driven Test Dispatcher" in result.stdout
    assert "--doc" in result.stdout
    assert "--output-report" in result.stdout


@pytest.mark.e2e
def test_cli_execute_document_with_output_report(
    sample_docx_spec: Path, tmp_path: Path, mock_base_url: str
) -> None:
    """Positive test: executes document via CLI main entrypoint with report output."""
    output_pdf = tmp_path / "cli_report.pdf"
    screenshots_dir = tmp_path / "cli_screenshots"

    exit_code = main(
        [
            "--doc",
            str(sample_docx_spec),
            "--output-report",
            str(output_pdf),
            "--screenshots-dir",
            str(screenshots_dir),
            "--base-url",
            mock_base_url,
            "--default-timeout",
            "5000",
            "--headless",
        ]
    )

    assert exit_code == 0
    assert output_pdf.exists()
    assert output_pdf.stat().st_size > 1024

    pdf_reader = pypdf.PdfReader(str(output_pdf))
    assert len(pdf_reader.pages) >= 1
    extracted_text = "".join(page.extract_text() or "" for page in pdf_reader.pages)
    assert "IT-001" in extracted_text
    assert "PASSED" in extracted_text


@pytest.mark.smoke
def test_cli_missing_required_doc_flag() -> None:
    """Negative test: running CLI without required --doc flag fails with error code 2."""
    result = subprocess.run(
        [sys.executable, "-m", "runners.doc_runner"],
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 2
    assert "the following arguments are required: --doc" in result.stderr


@pytest.mark.smoke
def test_cli_nonexistent_document_file(tmp_path: Path) -> None:
    """Negative test: running CLI pointing to missing file fails with non-zero exit code."""
    non_existent = tmp_path / "missing_file.docx"
    result = subprocess.run(
        [sys.executable, "-m", "runners.doc_runner", "--doc", str(non_existent)],
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode != 0
    assert "PackageNotFoundError" in result.stderr or "Error" in result.stderr
