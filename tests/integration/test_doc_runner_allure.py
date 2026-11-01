"""Integration tests for DocRunner Allure reporting and CLI integration."""

from __future__ import annotations

import json
from pathlib import Path

import docx
import pytest

from parsers.models import StepDefinition, TestCase
from runners.doc_runner import BlockedStepError, DocRunner, main

MOCK_HTML_PAGE = """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Allure Integration Mock Portal</title>
</head>
<body>
    <h1>Allure Reporting Portal Test Page</h1>
    <a href="#portal" id="portal-link">portal</a>
    <button id="action-btn" type="button">action</button>
</body>
</html>
"""


@pytest.fixture
def mock_html_path(tmp_path: Path) -> Path:
    """Create local HTML file for isolated Playwright execution."""
    html_file = tmp_path / "mock_allure_portal.html"
    html_file.write_text(MOCK_HTML_PAGE, encoding="utf-8")
    return html_file


@pytest.fixture
def mock_portal_url(mock_html_path: Path) -> str:
    """Return file:// URI for mock HTML page."""
    return mock_html_path.resolve().as_uri()


def _create_test_docx(
    file_path: Path,
    base_url: str,
    case_id: str = "ALLURE-001",
    fail_step: bool = False,
) -> Path:
    """Generate a minimal valid specification document."""
    doc = docx.Document()
    table = doc.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "Identificador del caso"
    table.rows[0].cells[1].text = case_id
    table.rows[0].cells[2].text = case_id

    table.rows[1].cells[0].text = "Módulo"
    table.rows[1].cells[1].text = "AllureReporting"
    table.rows[1].cells[2].text = "AllureReporting"

    table.rows[2].cells[0].text = "Ambiente de prueba"
    table.rows[2].cells[1].text = "IntegrationTest"
    table.rows[2].cells[2].text = "IntegrationTest"

    doc.add_paragraph(f"Link: {base_url}")
    doc.add_paragraph(f"Caso # 1: Execution for {case_id}")
    doc.add_paragraph("Objetivo de la Prueba")
    doc.add_paragraph("Verify Allure JSON lifecycle reporting in DocRunner.")
    doc.add_paragraph("Usuario: test_user")
    doc.add_paragraph("Clave: test_pass")

    doc.add_paragraph(
        'Step 1: (10/12/2026 10:00:00 AM) User left click on "portal (link)"'
    )
    if fail_step:
        doc.add_paragraph(
            'Step 2: (10/12/2026 10:00:02 AM) User left click on "#non-existent-button-xyz (button)"'
        )
    else:
        doc.add_paragraph(
            'Step 2: (10/12/2026 10:00:02 AM) User left click on "action (button)"'
        )

    doc.save(str(file_path))
    return file_path


@pytest.mark.e2e
def test_doc_runner_with_allure_dir_creates_result_and_container_positive(
    tmp_path: Path, mock_portal_url: str
) -> None:
    """Positive test: executes document with allure_dir and verifies Allure JSON files."""
    doc_path = _create_test_docx(tmp_path / "valid_spec.docx", mock_portal_url)
    allure_dir = tmp_path / "allure_output"
    screenshots_dir = tmp_path / "screenshots"

    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=5000,
        stop_on_failure=True,
        allure_dir=allure_dir,
    )
    report = runner.run_document(doc_path=doc_path, base_url=mock_portal_url)

    assert report.status == "PASSED"
    assert allure_dir.is_dir()

    result_files = list(allure_dir.glob("*-result.json"))
    container_files = list(allure_dir.glob("*-container.json"))

    assert len(result_files) == 1
    assert len(container_files) == 1

    result_data = json.loads(result_files[0].read_text(encoding="utf-8"))
    container_data = json.loads(container_files[0].read_text(encoding="utf-8"))

    assert result_data["status"] == "passed"
    assert "Execution for ALLURE-001" in result_data["name"]
    assert len(result_data["steps"]) == 2
    assert result_data["uuid"] in container_data["children"]


@pytest.mark.e2e
def test_doc_runner_without_allure_dir_negative(
    tmp_path: Path, mock_portal_url: str
) -> None:
    """Negative test: executes document without allure_dir and ensures no Allure output is generated."""
    doc_path = _create_test_docx(tmp_path / "no_allure_spec.docx", mock_portal_url)
    allure_dir = tmp_path / "should_not_exist_allure"
    screenshots_dir = tmp_path / "screenshots"

    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=5000,
        stop_on_failure=True,
    )
    report = runner.run_document(doc_path=doc_path, base_url=mock_portal_url)

    assert report.status == "PASSED"
    assert not allure_dir.exists()


@pytest.mark.e2e
def test_doc_runner_dual_reporting_pdf_and_allure_positive(
    tmp_path: Path, mock_portal_url: str
) -> None:
    """Positive test: executes document producing both PDF and Allure report artifacts."""
    doc_path = _create_test_docx(tmp_path / "dual_spec.docx", mock_portal_url)
    pdf_path = tmp_path / "executive_report.pdf"
    allure_dir = tmp_path / "allure_dual_output"
    screenshots_dir = tmp_path / "screenshots"

    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=5000,
        stop_on_failure=True,
    )
    report = runner.run_document(
        doc_path=doc_path,
        output_report_path=pdf_path,
        allure_dir=allure_dir,
        base_url=mock_portal_url,
    )

    assert report.status == "PASSED"
    assert pdf_path.is_file()
    assert pdf_path.stat().st_size > 1024

    assert allure_dir.is_dir()
    result_files = list(allure_dir.glob("*-result.json"))
    container_files = list(allure_dir.glob("*-container.json"))
    assert len(result_files) == 1
    assert len(container_files) == 1


@pytest.mark.e2e
def test_doc_runner_only_output_report_path_no_allure(
    tmp_path: Path, mock_portal_url: str
) -> None:
    """Positive & boundary test: only PDF path provided, no Allure directory or files created."""
    doc_path = _create_test_docx(tmp_path / "only_pdf_spec.docx", mock_portal_url)
    pdf_path = tmp_path / "only_pdf_report.pdf"
    screenshots_dir = tmp_path / "screenshots"

    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=5000,
        stop_on_failure=True,
    )
    report = runner.run_document(
        doc_path=doc_path,
        output_report_path=pdf_path,
        base_url=mock_portal_url,
    )

    assert report.status == "PASSED"
    assert pdf_path.is_file()
    assert pdf_path.stat().st_size > 1024

    # Verify no accidental allure directories created in tmp_path
    allure_json_files = list(tmp_path.glob("**/*-result.json"))
    assert len(allure_json_files) == 0


@pytest.mark.e2e
def test_doc_runner_only_allure_dir_no_pdf(
    tmp_path: Path, mock_portal_url: str
) -> None:
    """Positive & boundary test: only allure_dir provided, no PDF report is created."""
    doc_path = _create_test_docx(tmp_path / "only_allure_spec.docx", mock_portal_url)
    pdf_path = tmp_path / "uncreated_report.pdf"
    allure_dir = tmp_path / "only_allure_results"
    screenshots_dir = tmp_path / "screenshots"

    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=5000,
        stop_on_failure=True,
    )
    report = runner.run_document(
        doc_path=doc_path,
        output_report_path=None,
        allure_dir=allure_dir,
        base_url=mock_portal_url,
    )

    assert report.status == "PASSED"
    assert not pdf_path.exists()
    assert allure_dir.is_dir()
    result_files = list(allure_dir.glob("*-result.json"))
    assert len(result_files) == 1


@pytest.mark.smoke
def test_cli_allure_dir_argument_positive(tmp_path: Path, mock_portal_url: str) -> None:
    """Positive test: CLI main() accepts --allure-dir and generates Allure report files."""
    doc_path = _create_test_docx(tmp_path / "cli_spec.docx", mock_portal_url)
    allure_dir = tmp_path / "cli_allure_out"
    screenshots_dir = tmp_path / "cli_screenshots"

    exit_code = main(
        [
            "--doc",
            str(doc_path),
            "--allure-dir",
            str(allure_dir),
            "--screenshots-dir",
            str(screenshots_dir),
            "--base-url",
            mock_portal_url,
            "--default-timeout",
            "5000",
            "--headless",
        ]
    )

    assert exit_code == 0
    assert allure_dir.is_dir()
    result_files = list(allure_dir.glob("*-result.json"))
    assert len(result_files) == 1


@pytest.mark.smoke
def test_cli_missing_required_doc_with_allure_dir_negative(tmp_path: Path) -> None:
    """Negative test: CLI fails when --allure-dir is provided without required --doc flag."""
    allure_dir = tmp_path / "cli_allure_fail"
    with pytest.raises(SystemExit) as exc_info:
        main(["--allure-dir", str(allure_dir)])

    assert exc_info.value.code == 2


@pytest.mark.e2e
def test_step_failure_verdict_propagates_to_allure_positive(
    tmp_path: Path, mock_portal_url: str
) -> None:
    """Positive test: step failure status ('FAILED') propagates to Allure step and test status."""
    doc_path = _create_test_docx(
        tmp_path / "failed_spec.docx",
        mock_portal_url,
        case_id="ALLURE-FAIL-001",
        fail_step=True,
    )
    allure_dir = tmp_path / "allure_fail_output"
    screenshots_dir = tmp_path / "screenshots"

    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=1000,
        stop_on_failure=True,
        allure_dir=allure_dir,
    )
    report = runner.run_document(doc_path=doc_path, base_url=mock_portal_url)

    assert report.status == "FAILED"
    result_files = list(allure_dir.glob("*-result.json"))
    assert len(result_files) == 1

    result_data = json.loads(result_files[0].read_text(encoding="utf-8"))
    assert result_data["status"] == "failed"
    assert len(result_data["steps"]) == 2
    assert result_data["steps"][0]["status"] == "passed"
    assert result_data["steps"][1]["status"] == "failed"
    assert "statusDetails" in result_data["steps"][1]
    assert "TimeoutError" in result_data["steps"][1]["statusDetails"]["message"]


@pytest.mark.e2e
def test_step_blocked_verdict_propagates_to_allure_positive(
    tmp_path: Path, mock_portal_url: str
) -> None:
    """Positive test: blocked step status propagates to Allure Status.BROKEN ('broken')."""
    allure_dir = tmp_path / "allure_blocked_output"
    screenshots_dir = tmp_path / "screenshots"

    test_case = TestCase(
        case_id="ALLURE-BLOCK-001",
        module="VerificationModule",
        title="Blocked Test Scenario",
        objective="Verify blocked step propagation to Allure broken status.",
        credentials={},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Navigate to mock portal",
                action="navigate",
                target=mock_portal_url,
            ),
            StepDefinition(
                step_num=2,
                raw_text="Blocked prerequisite step",
                action="blocked",
                target="External verification dependency",
            ),
        ],
        environment="Integration",
        url=mock_portal_url,
    )

    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=2000,
        stop_on_failure=True,
        allure_dir=allure_dir,
    )
    report = runner.run_test_case(test_case=test_case, base_url=mock_portal_url)

    assert report.status == "BLOCKED"
    result_files = list(allure_dir.glob("*-result.json"))
    assert len(result_files) == 1

    result_data = json.loads(result_files[0].read_text(encoding="utf-8"))
    assert result_data["status"] == "broken"
    assert len(result_data["steps"]) == 2
    assert result_data["steps"][0]["status"] == "passed"
    assert result_data["steps"][1]["status"] == "broken"


@pytest.mark.e2e
def test_step_blocked_exception_propagates_to_allure_positive(
    tmp_path: Path, mock_portal_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Positive test: raising BlockedStepError in step execution propagates to Allure broken status."""
    allure_dir = tmp_path / "allure_blocked_exc_output"
    screenshots_dir = tmp_path / "screenshots"

    test_case = TestCase(
        case_id="ALLURE-BLOCK-002",
        module="VerificationModule",
        title="Blocked Exception Scenario",
        objective="Verify BlockedStepError propagation to Allure broken status.",
        credentials={},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Navigate to mock portal",
                action="navigate",
                target=mock_portal_url,
            ),
            StepDefinition(
                step_num=2,
                raw_text="Attempt click requiring external resource",
                action="click",
                target="#action-btn",
            ),
        ],
        environment="Integration",
        url=mock_portal_url,
    )

    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=2000,
        stop_on_failure=True,
        allure_dir=allure_dir,
    )

    def mock_dispatch(page, step, tc, base_url):
        if step.step_num == 2:
            raise BlockedStepError("External authentication service is unavailable")
        DocRunner._dispatch_step(runner, page, step, tc, base_url)

    monkeypatch.setattr(runner, "_dispatch_step", mock_dispatch)
    report = runner.run_test_case(test_case=test_case, base_url=mock_portal_url)

    assert report.status == "BLOCKED"
    result_files = list(allure_dir.glob("*-result.json"))
    assert len(result_files) == 1

    result_data = json.loads(result_files[0].read_text(encoding="utf-8"))
    assert result_data["status"] == "broken"
    assert len(result_data["steps"]) == 2
    assert result_data["steps"][0]["status"] == "passed"
    assert result_data["steps"][1]["status"] == "broken"
    assert "BlockedStepError" in result_data["steps"][1]["statusDetails"]["message"]


@pytest.mark.smoke
def test_doc_runner_invalid_allure_dir_negative(tmp_path: Path) -> None:
    """Negative test: specifying an existing file path as allure_dir raises NotADirectoryError."""
    fake_file = tmp_path / "not_a_dir.txt"
    fake_file.write_text("dummy", encoding="utf-8")

    runner = DocRunner(allure_dir=fake_file)
    test_case = TestCase(
        case_id="INVALID-DIR",
        module="Module",
        title="Invalid Dir Case",
        objective="Verify error on invalid allure_dir.",
        credentials={},
        steps=[],
    )

    with pytest.raises(NotADirectoryError):
        runner.run_test_case(test_case=test_case)


@pytest.mark.e2e
def test_doc_runner_blocked_step_masks_credentials_positive(
    tmp_path: Path, mock_portal_url: str
) -> None:
    """Positive test: blocked step error message redacts credentials present in step.target."""
    allure_dir = tmp_path / "allure_blocked_cred"
    screenshots_dir = tmp_path / "screenshots"

    secret_in_target = "BlockedSecretPass99"
    test_case = TestCase(
        case_id="BLOCK-CRED-001",
        module="VerificationModule",
        title="Blocked Credential Scenario",
        objective="Verify credential redaction in blocked step message.",
        credentials={"usuario": "admin", "clave": secret_in_target},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Blocked service invocation",
                action="blocked",
                target=f"https://admin:{secret_in_target}@api.corp.internal/auth",
            ),
        ],
        environment="Integration",
        url=mock_portal_url,
    )

    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        allure_dir=allure_dir,
    )
    report = runner.run_test_case(test_case=test_case, base_url=mock_portal_url)

    assert report.status == "BLOCKED"
    step_res = report.steps[0]
    assert step_res.error_message is not None
    assert secret_in_target not in step_res.error_message
    assert "[REDACTED]" in step_res.error_message


def test_doc_runner_blocked_step_clean_negative(
    tmp_path: Path, mock_portal_url: str
) -> None:
    """Negative test: clean blocked step target without credentials preserves message text."""
    screenshots_dir = tmp_path / "screenshots"
    test_case = TestCase(
        case_id="BLOCK-CLEAN-001",
        module="VerificationModule",
        title="Blocked Clean Scenario",
        objective="Verify clean blocked step does not over-redact text.",
        credentials={},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Blocked by routine maintenance",
                action="blocked",
                target="Routine maintenance window in progress",
            ),
        ],
        environment="Integration",
        url=mock_portal_url,
    )

    runner = DocRunner(headless=True, screenshots_dir=screenshots_dir)
    report = runner.run_test_case(test_case=test_case, base_url=mock_portal_url)

    assert report.status == "BLOCKED"
    step_res = report.steps[0]
    assert (
        step_res.error_message
        == "Step execution blocked: Routine maintenance window in progress"
    )


@pytest.mark.e2e
def test_full_dual_reporting_regression(tmp_path: Path, mock_portal_url: str) -> None:
    """T3.1 Regression test: runs dual PDF and Allure reporting, asserting schema compliance and zero secret leaks."""
    pdf_path = tmp_path / "regression_dual_report.pdf"
    allure_dir = tmp_path / "regression_allure_results"
    screenshots_dir = tmp_path / "regression_screenshots"

    secret_pass = "SuperSecretDualKey123!"
    secret_token = "eyJhbGciOiJIUzI1NiJ9.eyJ1c2VyIjoiZHVhbCJ9.dualSignature12345678"
    secret_email = "dual_auditor@security.corp"

    # Build spec document with credentials and sensitive step targets
    doc_path = tmp_path / "dual_regression_spec.docx"
    doc = docx.Document()
    table = doc.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "Identificador del caso"
    table.rows[0].cells[1].text = "REG-DUAL-001"
    table.rows[0].cells[2].text = "REG-DUAL-001"
    table.rows[1].cells[0].text = "Módulo"
    table.rows[1].cells[1].text = "SecurityReporting"
    table.rows[1].cells[2].text = "SecurityReporting"
    table.rows[2].cells[0].text = "Ambiente de prueba"
    table.rows[2].cells[1].text = "DualRegression"
    table.rows[2].cells[2].text = "DualRegression"

    doc.add_paragraph(f"Link: {mock_portal_url}")
    doc.add_paragraph(f'Caso # 1: Dual Execution with password="{secret_pass}"')
    doc.add_paragraph("Objetivo de la Prueba")
    doc.add_paragraph(
        f"Verify zero leakage for {secret_email} and token={secret_token}"
    )
    doc.add_paragraph("Usuario: dual_qa_user")
    doc.add_paragraph(f"Clave: {secret_pass}")

    doc.add_paragraph(
        'Step 1: (10/12/2026 10:00:00 AM) User left click on "portal (link)"'
    )
    doc.add_paragraph(
        'Step 2: (10/12/2026 10:00:02 AM) User left click on "action (button)"'
    )
    doc.save(str(doc_path))

    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=5000,
        stop_on_failure=True,
    )
    report = runner.run_document(
        doc_path=doc_path,
        output_report_path=pdf_path,
        allure_dir=allure_dir,
        base_url=mock_portal_url,
    )

    # 1. Assert runner execution and PDF output
    assert report.status == "PASSED"
    assert pdf_path.is_file()
    assert pdf_path.stat().st_size > 1024

    # 2. Assert Allure output files existence
    assert allure_dir.is_dir()
    result_files = list(allure_dir.glob("*-result.json"))
    container_files = list(allure_dir.glob("*-container.json"))
    assert len(result_files) == 1
    assert len(container_files) == 1

    result_raw = result_files[0].read_text(encoding="utf-8")
    container_raw = container_files[0].read_text(encoding="utf-8")
    result_data = json.loads(result_raw)
    container_data = json.loads(container_raw)

    # 3. Assert Allure JSON Schema Compliance
    for required_field in (
        "uuid",
        "name",
        "status",
        "start",
        "stop",
        "steps",
        "labels",
        "attachments",
    ):
        assert required_field in result_data, (
            f"Missing field '{required_field}' in Allure result"
        )

    for required_container_field in ("uuid", "name", "children", "start"):
        assert required_container_field in container_data, (
            f"Missing field '{required_container_field}' in Allure container"
        )

    assert result_data["uuid"] in container_data["children"]
    assert len(result_data["steps"]) == 2

    # 4. Assert ZERO Secret Leakage in emitted JSON
    for sensitive_value in (secret_pass, secret_token, secret_email):
        assert sensitive_value not in result_raw, (
            f"Secret value '{sensitive_value}' leaked in result JSON!"
        )
        assert sensitive_value not in container_raw, (
            f"Secret value '{sensitive_value}' leaked in container JSON!"
        )


@pytest.mark.e2e
def test_full_dual_reporting_regression_negative(tmp_path: Path) -> None:
    """Negative companion test: running non-existent document raises FileNotFoundError without leaving partial reports."""
    pdf_path = tmp_path / "non_existent_report.pdf"
    allure_dir = tmp_path / "non_existent_allure"
    missing_doc = tmp_path / "missing_file.docx"

    from docx.opc.exceptions import PackageNotFoundError

    runner = DocRunner(allure_dir=allure_dir)
    with pytest.raises((FileNotFoundError, PackageNotFoundError)):
        runner.run_document(doc_path=missing_doc, output_report_path=pdf_path)

    assert not pdf_path.exists()
    assert len(list(tmp_path.glob("*-result.json"))) == 0
