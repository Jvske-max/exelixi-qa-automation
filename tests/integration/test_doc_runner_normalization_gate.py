"""Integration and unit tests for DocRunner normalization gate, env var expansion, and manual steps (T5)."""

from __future__ import annotations

import base64
import re
from pathlib import Path

import docx
import pypdf
import pytest

from parsers.docx_parser import DocxParser
from parsers.models import StepDefinition, TestCase
from reporting.pdf_reporter import PDFReporter
from runners.doc_runner import DocRunner, _expand_env_vars, _mask_error_message, main

MOCK_HTML_CONTENT = """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Gate Test Portal</title>
</head>
<body>
    <h1>QA Portal Test Page</h1>
    <form id="login-form">
        <div>
            <label for="username">Username</label>
            <input type="text" id="username" name="username" placeholder="Username" />
        </div>
        <div>
            <label for="password">Password</label>
            <input type="password" id="password" name="password" placeholder="Password" />
        </div>
        <button type="submit" id="login-button">Login</button>
    </form>
    <div id="welcome-message" style="display:none;">Welcome to Dashboard</div>
    <script>
        document.getElementById('login-form').addEventListener('submit', function(e) {
            e.preventDefault();
            document.getElementById('welcome-message').style.display = 'block';
        });
    </script>
</body>
</html>
"""


@pytest.fixture
def mock_html_page(tmp_path: Path) -> Path:
    """Creates a local mock HTML page for isolated UI tests."""
    html_file = tmp_path / "mock_portal.html"
    html_file.write_text(MOCK_HTML_CONTENT, encoding="utf-8")
    return html_file


@pytest.fixture
def mock_base_url(mock_html_page: Path) -> str:
    """Returns local file URL for the mock HTML page."""
    return f"file://{mock_html_page.resolve()}"


@pytest.mark.smoke
def test_gate_exit_2_needs_info_without_browser(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture) -> None:
    """Pre-flight Gate: NEEDS_INFO exits with code 2, prints .gaps.md, never opens browser."""
    # 1. Setup docx and adjacent gaps.md
    doc_path = tmp_path / "spec_needs_info.docx"
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Identificador del caso"
    table.rows[0].cells[1].text = "GATE-001"
    table.rows[1].cells[0].text = "Estado de normalización"
    table.rows[1].cells[1].text = "NEEDS_INFO"
    doc.save(str(doc_path))

    gaps_path = tmp_path / "spec_needs_info.gaps.md"
    gaps_path.write_text("# Reporte de Faltantes\n- [BLOCKING] missing_url", encoding="utf-8")

    # 2. Monkeypatch sync_playwright to guarantee browser is never invoked
    def _fail_playwright(*args, **kwargs):
        pytest.fail("Browser / sync_playwright should NEVER be invoked for NEEDS_INFO status")

    monkeypatch.setattr("runners.doc_runner.sync_playwright", _fail_playwright)

    test_case = TestCase(
        case_id="GATE-001",
        module="Facturación",
        title="Caso bloqueado por faltantes",
        objective="Verificar detención temprana en gate",
        credentials={},
        steps=[
            StepDefinition(step_num=1, raw_text="Paso 1: clic en Boton", action="click", target="Boton")
        ],
        source_file=str(doc_path),
        normalization_status="NEEDS_INFO",
    )

    runner = DocRunner(headless=True, screenshots_dir=tmp_path / "screenshots")

    # Positive check: SystemExit with code 2 is raised
    with pytest.raises(SystemExit) as exc_info:
        runner.run_test_case(test_case)

    assert exc_info.value.code == 2

    # Verify gaps file path was printed to stdout
    captured = capsys.readouterr()
    assert str(gaps_path.resolve()) in captured.out or str(gaps_path) in captured.out

    # Negative check: run_document also stops with exit code 2
    with pytest.raises(SystemExit) as exc_doc_info:
        runner.run_document(doc_path)
    assert exc_doc_info.value.code == 2


@pytest.mark.smoke
def test_gate_exit_2_unsupported_without_browser(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture) -> None:
    """Pre-flight Gate: UNSUPPORTED exits with code 2, prints .gaps.md if present, never opens browser."""
    doc_path = tmp_path / "spec_unsupported.docx"
    doc = docx.Document()
    doc.save(str(doc_path))

    gaps_path = tmp_path / "spec_unsupported.gaps.md"
    gaps_path.write_text("# Gaps UNSUPPORTED", encoding="utf-8")

    def _fail_playwright(*args, **kwargs):
        pytest.fail("Browser / sync_playwright should NEVER be invoked for UNSUPPORTED status")

    monkeypatch.setattr("runners.doc_runner.sync_playwright", _fail_playwright)

    test_case = TestCase(
        case_id="GATE-002",
        module="OTP",
        title="Caso no automatizable",
        objective="Verificar que UNSUPPORTED no ejecuta pasos",
        credentials={},
        steps=[],
        source_file=str(doc_path),
        normalization_status="UNSUPPORTED",
    )

    runner = DocRunner(headless=True, screenshots_dir=tmp_path / "screenshots")

    with pytest.raises(SystemExit) as exc_info:
        runner.run_test_case(test_case)

    assert exc_info.value.code == 2

    captured = capsys.readouterr()
    assert str(gaps_path.resolve()) in captured.out or str(gaps_path) in captured.out


@pytest.mark.smoke
def test_missing_env_var_fails_without_leaking_secrets(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """Missing ${VAR} environment variable raises exception naming variable and redacting secrets."""
    # Ensure missing variable is NOT in environment
    monkeypatch.delenv("QA_USER_MISSING_VAR", raising=False)
    monkeypatch.setenv("QA_OTHER_SECRET", "SUPER_SECRET_VALUE_DO_NOT_LEAK_999")

    test_case = TestCase(
        case_id="GATE-003",
        module="Auth",
        title="Variable ausente en credenciales",
        objective="Verificar que falla nombrando la variable y nunca filtra secretos",
        credentials={
            "usuario": "${QA_USER_MISSING_VAR}",
            "clave": "${QA_OTHER_SECRET}",
        },
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Paso 1: escribir valor en campo",
                action="fill",
                target="username",
                value="algun_valor",
            )
        ],
        normalization_status="READY",
    )

    runner = DocRunner(headless=True, screenshots_dir=tmp_path / "screenshots")

    # Negative check 1: fails naming ${QA_USER_MISSING_VAR} without leaking secret
    with pytest.raises(ValueError) as exc_info:
        runner.run_test_case(test_case)

    err_msg = str(exc_info.value)
    assert "${QA_USER_MISSING_VAR}" in err_msg
    assert "SUPER_SECRET_VALUE_DO_NOT_LEAK_999" not in err_msg

    # Negative check 2: missing variable in step fill value also fails
    monkeypatch.setenv("QA_USER_MISSING_VAR", "valid_dummy_user")
    monkeypatch.delenv("QA_FILL_MISSING_VAL", raising=False)

    test_case_fill = TestCase(
        case_id="GATE-003B",
        module="Auth",
        title="Variable ausente en fill",
        objective="Verificar que falla nombrando variable en fill",
        credentials={"usuario": "dummy", "clave": "dummy"},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Paso 1: escribir valor en campo",
                action="fill",
                target="username",
                value="${QA_FILL_MISSING_VAL}",
            )
        ],
        normalization_status="READY",
    )

    with pytest.raises(ValueError) as exc_fill_info:
        runner.run_test_case(test_case_fill)

    err_fill_msg = str(exc_fill_info.value)
    assert "${QA_FILL_MISSING_VAL}" in err_fill_msg


@pytest.mark.e2e
def test_env_var_expansion_in_credentials_and_fill(
    tmp_path: Path, mock_base_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Positive test: ${VAR} references in credentials and fill values are properly expanded."""
    monkeypatch.setenv("QA_USER_TEST", "synthetic_user")
    monkeypatch.setenv("QA_PASS_TEST", "synthetic_pass")
    monkeypatch.setenv("QA_CODE_VAL", "FAC-998877")

    screenshots_dir = tmp_path / "screenshots_vars"
    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=5000,
    )

    test_case = TestCase(
        case_id="GATE-004",
        module="Facturación",
        title="Expansión de variables de entorno",
        objective="Validar sustitución exitosa de ${VAR}",
        url=mock_base_url,
        credentials={
            "username": "${QA_USER_TEST}",
            "password": "${QA_PASS_TEST}",
        },
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Navigate to mock portal",
                action="navigate",
                target=mock_base_url,
            ),
            StepDefinition(
                step_num=2,
                raw_text="Enter expanded username",
                action="fill",
                target="input[name='username']",
                value="${QA_CODE_VAL}",
            ),
            StepDefinition(
                step_num=3,
                raw_text="Enter password",
                action="fill",
                target="input[name='password']",
                value="${QA_PASS_TEST}",
            ),
            StepDefinition(
                step_num=4,
                raw_text="Click login",
                action="click",
                target="button#login-button",
            ),
            StepDefinition(
                step_num=5,
                raw_text="Verify greeting",
                action="verify",
                target="Welcome to Dashboard",
            ),
        ],
        normalization_status="READY",
    )

    report = runner.run_test_case(test_case=test_case, base_url=mock_base_url)

    assert report.status == "PASSED"
    assert report.passed_count == 5
    assert report.failed_count == 0
    # Credentials were expanded
    assert test_case.credentials["username"] == "synthetic_user"
    assert test_case.credentials["password"] == "synthetic_pass"


@pytest.mark.e2e
def test_partial_status_with_manual_steps(tmp_path: Path, mock_base_url: str) -> None:
    """All executable steps pass + manual_steps present -> global status PARTIAL and CLI returns 0."""
    screenshots_dir = tmp_path / "screenshots_partial"
    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=5000,
    )

    test_case = TestCase(
        case_id="GATE-005",
        module="Cobranzas",
        title="Caso mixto con pasos manuales",
        objective="Verificar status PARTIAL y filas MANUAL",
        url=mock_base_url,
        credentials={},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Navegar al portal",
                action="navigate",
                target=mock_base_url,
            ),
            StepDefinition(
                step_num=2,
                raw_text="Escribir usuario",
                action="fill",
                target="input[name='username']",
                value="operador",
            ),
        ],
        manual_steps=[
            "Adjuntar certificado digital fiscal desde token USB",
            "Verificar código OTP recibido por SMS",
        ],
        normalization_status="READY",
    )

    report = runner.run_test_case(test_case=test_case, base_url=mock_base_url)

    # 1. Verify global status is PARTIAL
    assert report.status == "PARTIAL"
    assert report.passed_count == 2
    assert report.failed_count == 0

    # 2. Verify manual steps are present with action="manual" and status="MANUAL"
    manual_results = [s for s in report.steps if s.action == "manual"]
    assert len(manual_results) == 2
    assert all(s.status == "MANUAL" for s in manual_results)
    assert manual_results[0].target == "Adjuntar certificado digital fiscal desde token USB"
    assert manual_results[1].target == "Verificar código OTP recibido por SMS"

    # 3. PDF reporter supports PARTIAL badge and MANUAL steps
    output_pdf = tmp_path / "report_partial.pdf"
    reporter = PDFReporter()
    generated_pdf = reporter.generate_report(report, str(output_pdf))
    assert Path(generated_pdf).exists()

    pdf_reader = pypdf.PdfReader(generated_pdf)
    text = "".join(p.extract_text() or "" for p in pdf_reader.pages)
    assert "PARTIAL" in text
    assert "GATE-005" in text
    assert "MANUAL" in text

    # 4. CLI exits with code 0 for PARTIAL status
    # Create a docx file that produces this test case
    docx_file = tmp_path / "spec_partial.docx"
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=3)
    t.rows[0].cells[0].text = "Identificador del caso"
    t.rows[0].cells[1].text = "GATE-005"
    t.rows[0].cells[2].text = "GATE-005"
    t.rows[1].cells[0].text = "Estado de normalización"
    t.rows[1].cells[1].text = "READY"
    t.rows[1].cells[2].text = "READY"

    doc.add_paragraph(f"Link: {mock_base_url}")
    doc.add_paragraph("Caso 1: Flujo parcial")
    doc.add_paragraph("Objetivo: Test parcial")
    doc.add_paragraph("Paso 1: navegar a " + mock_base_url)
    doc.add_paragraph("Paso 2: escribir operador en Username")
    doc.add_paragraph("Pendiente manual 1: Adjuntar certificado digital fiscal desde token USB")
    doc.add_paragraph("Pendiente manual 2: Verificar código OTP recibido por SMS")
    doc.save(str(docx_file))

    cli_exit_code = main(
        [
            "--doc",
            str(docx_file),
            "--screenshots-dir",
            str(screenshots_dir),
            "--base-url",
            mock_base_url,
            "--headless",
        ]
    )
    assert cli_exit_code == 0


@pytest.mark.smoke
def test_docx_parser_extracts_normalization_status_and_manual_steps(tmp_path: Path) -> None:
    """DocxParser reads 'Estado de normalización' row and 'Pendiente manual' paragraphs."""
    docx_file = tmp_path / "parser_test.docx"
    doc = docx.Document()

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Identificador del caso"
    table.rows[0].cells[1].text = "TEST-PARSER-01"
    table.rows[1].cells[0].text = "Estado de normalización"
    table.rows[1].cells[1].text = "READY"

    doc.add_paragraph("Link: https://example.test")
    doc.add_paragraph("Caso 1: Normalizer parsing check")
    doc.add_paragraph("Objetivo de la Prueba: Test parser additions")
    doc.add_paragraph("Paso 1: clic en Iniciar")
    doc.add_paragraph("Pendiente manual 1: Subir archivo PDF al servidor")
    doc.add_paragraph("Pendiente manual 2: Validar SMS de confirmación")
    doc.save(str(docx_file))

    parsed = DocxParser.parse(docx_file)

    # Positive checks
    assert parsed.normalization_status == "READY"
    assert len(parsed.manual_steps) == 2
    assert parsed.manual_steps[0] == "Subir archivo PDF al servidor"
    assert parsed.manual_steps[1] == "Validar SMS de confirmación"
    assert len(parsed.steps) == 1
    assert parsed.steps[0].action == "click"

    # Negative check: legacy document without normalization table row
    legacy_file = tmp_path / "legacy_parser_test.docx"
    doc_leg = docx.Document()
    t_leg = doc_leg.add_table(rows=1, cols=2)
    t_leg.rows[0].cells[0].text = "Identificador del caso"
    t_leg.rows[0].cells[1].text = "LEGACY-01"
    doc_leg.add_paragraph("Step 1: (10/12/2026) User left click on \"Boton\"")
    doc_leg.save(str(legacy_file))

    parsed_leg = DocxParser.parse(legacy_file)
    assert parsed_leg.normalization_status is None
    assert parsed_leg.manual_steps == []
    assert len(parsed_leg.steps) == 1


@pytest.mark.e2e
def test_legacy_format_backward_compatibility(tmp_path: Path, mock_base_url: str) -> None:
    """Documents without normalization_status (legacy) retain previous fallback behavior and pass."""
    screenshots_dir = tmp_path / "screenshots_legacy"
    runner = DocRunner(
        headless=True,
        screenshots_dir=screenshots_dir,
        default_timeout=5000,
    )

    # Legacy TestCase: normalization_status is None
    legacy_case = TestCase(
        case_id="LEGACY-001",
        module="LegacyModule",
        title="Escenario PSR sin normalizar",
        objective="Verificar compatibilidad 100% aditiva con formato anterior",
        url=mock_base_url,
        credentials={},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Navigate to mock portal",
                action="navigate",
                target=mock_base_url,
            ),
            StepDefinition(
                step_num=2,
                raw_text="Click login button",
                action="click",
                target="button#login-button",
            ),
        ],
        normalization_status=None,
    )

    report = runner.run_test_case(test_case=legacy_case, base_url=mock_base_url)

    assert report.status == "PASSED"
    assert report.passed_count == 2
    assert report.failed_count == 0
    assert len([s for s in report.steps if s.action == "manual"]) == 0


@pytest.mark.smoke
def test_doc_runner_masks_secrets_in_step_error_message(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify that DocRunner masks step.value and credential patterns in StepResult.error_message on failure."""
    secret_step_val = "SECRET_STEP_VALUE_XYZ"
    secret_cred_val = "SECRET_CRED_VALUE_123"

    runner = DocRunner(headless=True, screenshots_dir=tmp_path / "screenshots")

    def _failing_dispatch(page, step, test_case, base_url):
        raise RuntimeError(
            f"Failed to process target {step.target} with value {step.value} and password: {secret_cred_val}"
        )

    monkeypatch.setattr(runner, "_dispatch_step", _failing_dispatch)

    test_case = TestCase(
        case_id="GATE-MASK-001",
        module="Auth",
        title="Error message secret masking test",
        objective="Verificar enmascaramiento de secretos en error_message",
        credentials={"password": secret_cred_val},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Escribir clave secreta",
                action="fill",
                target="input#pass",
                value=secret_step_val,
            )
        ],
        normalization_status="READY",
    )

    report = runner.run_test_case(test_case=test_case)

    assert report.status == "FAILED"
    assert len(report.steps) == 1
    err_msg = report.steps[0].error_message
    assert err_msg is not None
    # Verify sensitive data is NOT present in error_message
    assert secret_step_val not in err_msg
    assert secret_cred_val not in err_msg
    # Verify masking indicator is present
    assert "[REDACTED]" in err_msg or "***" in err_msg


@pytest.mark.smoke
def test_doc_runner_legacy_no_hardcoded_credentials_or_corporate_urls(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify doc_runner.py has no hardcoded corporate URLs or credentials in legacy branch."""
    doc_runner_source = Path("runners/doc_runner.py").read_text(encoding="utf-8")

    # 1. Static inspection: verify corporate domain and default personal credentials are not hardcoded
    assert not re.search(r"[a-z0-9.-]*seguros\.com", doc_runner_source, re.IGNORECASE), (
        "runners/doc_runner.py must not contain hardcoded corporate URLs"
    )
    assert not re.search(r"\brmu[ñn]oz\b", doc_runner_source, re.IGNORECASE), (
        "runners/doc_runner.py must not contain hardcoded personal credentials"
    )

    # 2. Dynamic inspection: legacy fallback resolves to empty strings when env vars are unset
    monkeypatch.delenv("QA_BASE_URL", raising=False)
    monkeypatch.delenv("QA_USERNAME", raising=False)
    monkeypatch.delenv("QA_PASSWORD", raising=False)

    legacy_case = TestCase(
        case_id="LEGACY-CHECK-01",
        module="LegacyCheck",
        title="Verificación de fallbacks sin env vars",
        objective="Garantizar fallbacks vacíos o impersonales",
        credentials={},
        steps=[
            StepDefinition(step_num=1, raw_text="Dummy", action="navigate", target="")
        ],
        normalization_status=None,
    )

    runner = DocRunner(headless=True)
    dispatched_url = []

    def _mock_dispatch(page, step, tc, resolved_base_url):
        dispatched_url.append(resolved_base_url)

    monkeypatch.setattr(runner, "_dispatch_step", _mock_dispatch)
    runner.run_test_case(legacy_case)
    assert dispatched_url == [""]

    # Verify login credentials fallback resolves to empty strings when env vars are unset
    captured_creds = {}

    class DummyLoginPage:
        def __init__(self, page):
            pass

        def login(self, username, password):
            captured_creds["username"] = username
            captured_creds["password"] = password

    monkeypatch.setattr("runners.doc_runner.LoginPage", DummyLoginPage)

    login_step_route = StepDefinition(
        step_num=1,
        raw_text="Login",
        action="fill",
        target="usuario",
        value="",
    )
    test_runner = DocRunner(headless=True)
    test_runner._dispatch_step(None, login_step_route, legacy_case, "")
    # Check captured creds from login branch
    assert captured_creds.get("username") == ""
    assert captured_creds.get("password") == ""

    # Negative check: when env vars are explicitly provided, they are resolved properly
    monkeypatch.setenv("QA_BASE_URL", "https://custom.example.test")
    monkeypatch.setenv("QA_USERNAME", "custom_user")
    monkeypatch.setenv("QA_PASSWORD", "custom_pass")

    dispatched_custom = []

    def _mock_dispatch_custom(page, step, tc, resolved_base_url):
        dispatched_custom.append(resolved_base_url)

    monkeypatch.setattr(runner, "_dispatch_step", _mock_dispatch_custom)
    runner.run_test_case(legacy_case)
    assert dispatched_custom == ["https://custom.example.test"]

    captured_creds.clear()
    test_runner._dispatch_step(None, login_step_route, legacy_case, "")
    assert captured_creds.get("username") == "custom_user"
    assert captured_creds.get("password") == "custom_pass"


@pytest.mark.smoke
def test_doc_runner_normalized_navigate_missing_url_raises_value_error(monkeypatch: pytest.MonkeyPatch) -> None:
    """Normalized test case dispatching navigate with empty target_url and base_url raises ValueError."""
    monkeypatch.delenv("QA_BASE_URL", raising=False)

    normalized_case = TestCase(
        case_id="NORM-NAV-01",
        module="Auth",
        title="Navegación sin URL",
        objective="Verificar fallo por URL faltante en caso normalizado",
        credentials={},
        steps=[
            StepDefinition(step_num=1, raw_text="Navigate", action="navigate", target="")
        ],
        url=None,
        normalization_status="READY",
    )

    runner = DocRunner(headless=True)
    step = normalized_case.steps[0]

    # Negative check: dispatch raises ValueError when URL is missing
    with pytest.raises(ValueError, match="Missing target URL for navigation step in normalized test case"):
        runner._dispatch_step(None, step, normalized_case, "")

    # Positive check: valid URL does not raise Missing target URL error
    class MockLoginPage:
        def __init__(self, page):
            self.navigated = None
        def navigate(self, url):
            self.navigated = url

    monkeypatch.setattr("runners.doc_runner.LoginPage", MockLoginPage)
    valid_step = StepDefinition(step_num=1, raw_text="Navigate", action="navigate", target="https://example.test/app")
    runner._dispatch_step(None, valid_step, normalized_case, "")


@pytest.mark.smoke
def test_doc_runner_normalized_login_missing_credentials_raises_value_error(monkeypatch: pytest.MonkeyPatch) -> None:
    """Normalized test case dispatching login without username or password raises ValueError."""
    monkeypatch.delenv("QA_USERNAME", raising=False)
    monkeypatch.delenv("QA_PASSWORD", raising=False)

    normalized_case = TestCase(
        case_id="NORM-LOGIN-01",
        module="Auth",
        title="Login sin credenciales",
        objective="Verificar fallo por credenciales ausentes en caso normalizado",
        credentials={},
        steps=[
            StepDefinition(step_num=1, raw_text="Login al sistema", action="login", target="login")
        ],
        normalization_status="READY",
    )

    runner = DocRunner(headless=True)
    step = normalized_case.steps[0]

    # Negative check 1: completely missing credentials raises ValueError
    with pytest.raises(ValueError, match="Missing login credentials for normalized test case"):
        runner._dispatch_step(None, step, normalized_case, "")

    # Negative check 2: username present but password missing raises ValueError
    normalized_case.credentials = {"username": "admin_only"}
    with pytest.raises(ValueError, match="Missing login credentials for normalized test case"):
        runner._dispatch_step(None, step, normalized_case, "")

    # Positive check: both credentials provided dispatches successfully
    captured = {}
    class MockLoginPage:
        def __init__(self, page):
            pass
        def login(self, username, password):
            captured["username"] = username
            captured["password"] = password

    monkeypatch.setattr("runners.doc_runner.LoginPage", MockLoginPage)
    normalized_case.credentials = {"username": "valid_user", "password": "valid_password"}
    runner._dispatch_step(None, step, normalized_case, "")
    assert captured["username"] == "valid_user"
    assert captured["password"] == "valid_password"


@pytest.mark.smoke
def test_expand_env_vars_strict_uppercase_only(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify that only strict uppercase variables ${UPPER_CASE} are expanded, and lowercase ones are ignored."""
    monkeypatch.setenv("QA_USER", "real_qa_user")
    monkeypatch.setenv("qa_user", "unexpected_expansion")
    monkeypatch.setenv("Qa_User", "unexpected_mixed")

    # Positive check: strict uppercase variable is expanded properly
    assert _expand_env_vars("${QA_USER}") == "real_qa_user"
    assert _expand_env_vars("Prefix ${QA_USER} Suffix") == "Prefix real_qa_user Suffix"

    # Negative check 1: variable with lowercase is NOT expanded even if present in os.environ
    assert _expand_env_vars("${qa_user}") == "${qa_user}"
    assert _expand_env_vars("${Qa_User}") == "${Qa_User}"
    assert _expand_env_vars("User: ${qa_user}") == "User: ${qa_user}"

    # Negative check 2: unset lowercase variable does NOT raise ValueError
    monkeypatch.delenv("qa_user", raising=False)
    assert _expand_env_vars("${qa_user}") == "${qa_user}"


@pytest.mark.smoke
def test_doc_runner_normalized_generic_fill_missing_credentials_raises_value_error() -> None:
    """Normalized test case dispatching generic fill on user/password field without credentials raises ValueError."""
    class DummyLocator:
        first = None

        def __init__(self):
            self.first = self
            self.filled_val = None

        def count(self):
            return 1

        def fill(self, val, timeout=None):
            self.filled_val = val

    class DummyPage:
        def __init__(self):
            self.loc = DummyLocator()

        def locator(self, *args, **kwargs):
            return self.loc

        def get_by_placeholder(self, *args, **kwargs):
            return self.loc

        def get_by_label(self, *args, **kwargs):
            return self.loc

        def get_by_role(self, *args, **kwargs):
            return self.loc

    runner = DocRunner(headless=True)
    dummy_page = DummyPage()

    # Negative check 1: fill on username field without credentials in normalized case raises ValueError
    case_no_user = TestCase(
        case_id="NORM-FILL-01",
        module="Auth",
        title="Fill usuario sin credenciales",
        objective="Verificar fallo inmediato en fill genérico",
        credentials={},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Escribir en campo_usuario",
                action="fill",
                target="campo_usuario",
                value="",
            )
        ],
        normalization_status="READY",
    )
    with pytest.raises(ValueError, match="Missing credential value for normalized test case step"):
        runner._dispatch_step(dummy_page, case_no_user.steps[0], case_no_user, "")

    # Negative check 2: fill on password field without credentials in normalized case raises ValueError
    case_no_pass = TestCase(
        case_id="NORM-FILL-02",
        module="Auth",
        title="Fill clave sin credenciales",
        objective="Verificar fallo inmediato en fill genérico",
        credentials={},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Escribir en campo_clave",
                action="fill",
                target="campo_clave",
                value="",
            )
        ],
        normalization_status="READY",
    )
    with pytest.raises(ValueError, match="Missing credential value for normalized test case step"):
        runner._dispatch_step(dummy_page, case_no_pass.steps[0], case_no_pass, "")

    # Positive check 1: normalized case WITH credential resolves and fills without error
    case_valid = TestCase(
        case_id="NORM-FILL-03",
        module="Auth",
        title="Fill con credencial presente",
        objective="Verificar éxito cuando credencial está disponible",
        credentials={"username": "valid_user"},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Escribir en campo_usuario",
                action="fill",
                target="campo_usuario",
                value="",
            )
        ],
        normalization_status="READY",
    )
    runner._dispatch_step(dummy_page, case_valid.steps[0], case_valid, "")
    assert dummy_page.loc.filled_val == "valid_user"

    # Positive check 2: legacy case (normalization_status=None) does not raise ValueError
    case_legacy = TestCase(
        case_id="LEGACY-FILL-01",
        module="Auth",
        title="Legacy fill sin credenciales",
        objective="Compatibilidad hacia atrás para casos legacy",
        credentials={},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Escribir en campo_usuario",
                action="fill",
                target="campo_usuario",
                value="",
            )
        ],
        normalization_status=None,
    )
    runner._dispatch_step(dummy_page, case_legacy.steps[0], case_legacy, "")
    assert dummy_page.loc.filled_val == ""


@pytest.mark.smoke
def test_mask_error_message_masks_sensitive_target_content() -> None:
    """Verify that sensitive parameters and tokens embedded in step.target are masked in error messages."""
    step = StepDefinition(
        step_num=1,
        raw_text="Navegar a URL con token de autenticación",
        action="navigate",
        target="https://qa.example.test/auth?token=super_secret_token_12345",
        value="",
    )
    test_case = TestCase(
        case_id="SEC-TARGET-01",
        module="Auth",
        title="Sanitizar target sensible",
        objective="Verificar que parámetros y tokens sensibles en step.target sean enmascarados",
        credentials={},
        steps=[step],
    )
    raw_error = "Timeout 30000ms exceeded navigating to https://qa.example.test/auth?token=super_secret_token_12345"
    masked = _mask_error_message(raw_error, step, test_case)

    # Negative check 1: sensitive token must not appear in masked error message
    assert "super_secret_token_12345" not in masked

    # Positive check 1: token query parameter value is specifically replaced by [REDACTED]
    assert "https://qa.example.test/auth?token=[REDACTED]" in masked

    # Negative check 2: error referencing extracted sensitive token value directly is also masked
    raw_error_direct = "Navigation error: token value super_secret_token_12345 was rejected by endpoint"
    masked_direct = _mask_error_message(raw_error_direct, step, test_case)
    assert "super_secret_token_12345" not in masked_direct
    assert "[REDACTED]" in masked_direct

    # Positive check 2: non-sensitive URL targets are preserved without alterations
    normal_step = StepDefinition(
        step_num=2,
        raw_text="Navegar a panel principal",
        action="navigate",
        target="https://qa.example.test/portal/dashboard",
        value="",
    )
    normal_error = "Timeout 30000ms exceeded navigating to https://qa.example.test/portal/dashboard"
    masked_normal = _mask_error_message(normal_error, normal_step, test_case)
    assert "https://qa.example.test/portal/dashboard" in masked_normal


@pytest.mark.smoke
def test_doc_runner_normalized_ignores_ambient_env_var_fallbacks(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify that normalized test cases ignore ambient env var fallbacks for URL and credentials."""
    monkeypatch.setenv("QA_USERNAME", "env_user")
    monkeypatch.setenv("QA_PASSWORD", "env_pass")
    monkeypatch.setenv("QA_BASE_URL", "https://env.example.test")

    # 1. Negative case: normalized test case with empty credentials and url
    normalized_case = TestCase(
        case_id="NORM-GATE-01",
        module="Auth",
        title="Test sin credenciales declaradas",
        objective="Verificar que no usa fallbacks de entorno",
        credentials={},
        url="",
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Iniciar sesión en la aplicación",
                action="login",
                target="login",
                value="",
            )
        ],
        normalization_status="READY",
    )

    runner = DocRunner(headless=True)

    # Assert that _dispatch_step with route == 'login' raises ValueError indicating missing credentials
    with pytest.raises(ValueError, match="Missing login credentials for normalized test case"):
        runner._dispatch_step(page=None, step=normalized_case.steps[0], test_case=normalized_case, base_url="")

    # Assert that resolved_base_url in _execute_test_case does not take env var fallback
    monkeypatch.setattr(runner, "_capture_screenshot", lambda p, fn: None)
    dispatched_urls = []

    def _mock_dispatch(page, step, tc, resolved_base_url):
        dispatched_urls.append(resolved_base_url)

    monkeypatch.setattr(runner, "_dispatch_step", _mock_dispatch)
    runner._execute_test_case(test_case=normalized_case, page=None)
    assert dispatched_urls == [""], f"Expected empty resolved_base_url, got {dispatched_urls}"

    # 2. Positive case: legacy test case (normalization_status is None) still takes ambient fallbacks
    legacy_case = TestCase(
        case_id="LEGACY-01",
        module="Auth",
        title="Legacy case using env fallbacks",
        objective="Verificar retrocompatibilidad en legacy",
        credentials={},
        url="",
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Navegar a inicio",
                action="navigate",
                target="",
                value="",
            )
        ],
        normalization_status=None,
    )
    dispatched_legacy_urls = []
    monkeypatch.setattr(runner, "_dispatch_step", lambda p, s, tc, r_url: dispatched_legacy_urls.append(r_url))
    runner._execute_test_case(test_case=legacy_case, page=None)
    assert dispatched_legacy_urls == ["https://env.example.test"]


@pytest.mark.smoke
def test_doc_runner_normalized_generic_fill_contrasena_raises_value_error(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify that normalized cases with empty value on contraseña/contrasena raise ValueError."""
    runner = DocRunner(headless=True)

    # 1. Negative case: target contains 'Contraseña' with empty value and no credentials
    normalized_case = TestCase(
        case_id="NORM-FILL-01",
        module="Auth",
        title="Paso fill contraseña sin valor",
        objective="Verificar que levanta ValueError en contraseña",
        credentials={},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Ingresar contraseña en campo",
                action="fill",
                target="Campo Contraseña",
                value="",
            )
        ],
        normalization_status="READY",
    )

    with pytest.raises(ValueError, match="Missing credential value for normalized test case step"):
        runner._dispatch_generic_action(
            page=None,
            step=normalized_case.steps[0],
            test_case=normalized_case,
            base_url="",
        )

    # Negative case variant: target contains 'contrasena' without tilde
    step_contrasena = StepDefinition(
        step_num=2,
        raw_text="Ingresar contrasena en campo",
        action="fill",
        target="input_contrasena",
        value="",
    )
    with pytest.raises(ValueError, match="Missing credential value for normalized test case step"):
        runner._dispatch_generic_action(
            page=None,
            step=step_contrasena,
            test_case=normalized_case,
            base_url="",
        )

    # 2. Positive case: credentials provide the password or explicit value is supplied
    class DummyLocator:
        def __init__(self):
            self.filled_val = None

        def fill(self, val, timeout=None):
            self.filled_val = val

    dummy_locator = DummyLocator()
    monkeypatch.setattr(runner, "_locate_for_fill", lambda p, t: dummy_locator)

    runner._dispatch_generic_action(
        page=None,
        step=StepDefinition(
            step_num=3,
            raw_text="Ingresar contraseña válida",
            action="fill",
            target="Campo Contraseña",
            value="MiPasswordSeguro123",
        ),
        test_case=normalized_case,
        base_url="",
    )
    assert dummy_locator.filled_val == "MiPasswordSeguro123"


@pytest.mark.smoke
def test_mask_error_message_masks_pii_patterns() -> None:
    """Verify that _mask_error_message redacts cédulas, RIF, emails, and JWT tokens."""
    dummy_step = StepDefinition(
        step_num=1,
        raw_text="Paso genérico",
        action="click",
        target="Boton",
        value="",
    )
    dummy_case = TestCase(
        case_id="PII-MASK-01",
        module="Auth",
        title="Enmascarar PII",
        objective="Verificar redacción de datos sensibles en errores",
        credentials={},
        steps=[dummy_step],
    )

    # 1. Negative check: raw error containing PII must have all sensitive elements redacted
    fake_header = base64.urlsafe_b64encode(b'{"alg":"HS256","typ":"JWT"}').decode().rstrip("=")
    fake_payload = base64.urlsafe_b64encode(b'{"sub":"1234567890"}').decode().rstrip("=")
    fake_signature = base64.urlsafe_b64encode(b"dummy-synthetic-signature-bytes").decode().rstrip("=")
    synthetic_jwt = f"{fake_header}.{fake_payload}.{fake_signature}"

    raw_error = (
        f"Error in session: cedula V-12345678, email analista@test.com, "
        f"RIF J-12345678-9, and token {synthetic_jwt}"
    )

    masked = _mask_error_message(raw_error, dummy_step, dummy_case)

    assert "V-12345678" not in masked, "Cédula was not masked"
    assert "analista@test.com" not in masked, "Email was not masked"
    assert "J-12345678-9" not in masked, "RIF was not masked"
    assert synthetic_jwt not in masked, "JWT was not masked"
    assert masked.count("[REDACTED]") >= 4

    # 2. Positive check: regular technical error message without PII remains unaltered
    clean_error = "Timeout 5000ms exceeded while waiting for selector #submit-btn"
    masked_clean = _mask_error_message(clean_error, dummy_step, dummy_case)
    assert masked_clean == clean_error


@pytest.mark.smoke
def test_doc_runner_normalized_requires_base_url_for_interactive_steps(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Normalized test case executing interactive steps on about:blank without base_url raises ValueError."""

    class DummyPage:
        def __init__(self, initial_url: str = "about:blank"):
            self.url = initial_url
            self.navigated = None

        def goto(self, url: str, **kwargs):
            self.url = url
            self.navigated = url

    class MockLoginPage:
        def __init__(self, page):
            self.page = page

        def navigate(self, url: str):
            if self.page is not None:
                self.page.url = url
                self.page.navigated = url

        def login(self, username: str, password: str):
            pass

    monkeypatch.setattr("runners.doc_runner.LoginPage", MockLoginPage)

    normalized_case = TestCase(
        case_id="GATE-INT-001",
        module="Auth",
        title="Paso interactivo sin URL base",
        objective="Verificar que paso interactivo sobre about:blank exige base_url en caso normalizado",
        credentials={"usuario": "admin", "clave": "secret"},
        steps=[
            StepDefinition(
                step_num=1,
                raw_text="Paso 1: clic en Iniciar Sesión",
                action="click",
                target="Iniciar Sesión",
            )
        ],
        url="",
        normalization_status="READY",
    )

    runner = DocRunner(headless=True)
    dummy_page = DummyPage("about:blank")
    step = normalized_case.steps[0]

    # Negative check: running interactive step on about:blank with base_url="" raises ValueError
    with pytest.raises(
        ValueError, match="Missing base URL for normalized test case with interactive steps"
    ):
        runner._dispatch_step(dummy_page, step, normalized_case, base_url="")

    # Positive check: providing a valid base_url allows dispatching and navigates away from about:blank
    runner._dispatch_step(
        dummy_page, step, normalized_case, base_url="https://qa.example.test/portal"
    )
    assert dummy_page.url == "https://qa.example.test/portal"











